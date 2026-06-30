"""Build the Aegis-40 FER plain-language HANDBOOK / GLOSSARY (.docx).

A from-scratch, simple-words explainer of every term, symbol, unit and concept that
shows up in the Aegis-40 Final Evaluation Report and its figures (core physics, fuel,
cooling, the power cycle, the plant flow-diagram symbols, safety, waste, and all the
acronyms). Aimed at a teammate who is NOT a specialist in that sub-area.

Output: docs/competition/fer/Aegis40_FER_Handbook-Glossary.docx
Run:    py scripts/build_glossary_docx.py     (requires python-docx)
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUTDOC = os.path.join(ROOT, "docs", "competition", "fer",
                      "Aegis40_FER_Handbook-Glossary.docx")

INK = RGBColor(0x22, 0x22, 0x22)
ACCENT = RGBColor(0x1F, 0x4E, 0x79)     # deep blue
SHADE = "DDE3EA"

# ===========================================================================
# CONTENT  — each glossary entry is (term, symbol/unit or None, plain definition)
# ===========================================================================

INTRO = (
    "This handbook explains, in plain language, the words, symbols and units used in the "
    "Aegis-40 Final Evaluation Report (FER) and its diagrams. It is written for someone who is "
    "smart but is NOT a specialist in that particular sub-area — if at any point in the FER you "
    "thought \"wait, what does that actually mean?\", look it up here. Terms are grouped by topic; "
    "there is an A–Z acronym list and a units table near the end. Where a symbol or unit is "
    "commonly used, it is shown in [brackets] after the term."
)

SECTIONS = [
    ("1.  The Big Picture — What Kind of Reactor Is This?",
     "Aegis-40 is a small nuclear power plant. These terms describe its overall type and philosophy.",
     [
        ("Nuclear fission", None,
         "Splitting a heavy atom (uranium-235) by hitting it with a neutron. The split releases a "
         "lot of heat plus 2–3 new neutrons, which go on to split more atoms — a self-sustaining "
         "chain reaction. That heat is the entire energy source of the plant."),
        ("Reactor core", None,
         "The part where fission happens — the bundle of fuel where the chain reaction is kept going "
         "in a controlled way. Everything else in the plant exists to remove its heat, control it, "
         "or turn its heat into electricity."),
        ("PWR — Pressurised Water Reactor", None,
         "The most common reactor type. Ordinary water both cools the core and slows neutrons down "
         "(see 'moderator'). The water is kept under high pressure so it stays liquid even when very "
         "hot (it does NOT boil in the core)."),
        ("SMR — Small Modular Reactor", None,
         "A small reactor (here 40 MW of electricity, vs ~1000 MW for a big plant) built from "
         "factory-made modules instead of poured on-site piece by piece. Smaller = cheaper to build, "
         "safer to cool passively, and quicker to construct. Aegis-40 is an SMR."),
        ("iPWR — integral PWR", None,
         "A PWR where the whole primary system — core, pressuriser, steam generator, control-rod "
         "drives — lives INSIDE one single pressure vessel, instead of being connected by big "
         "external pipes. Fewer big pipes means fewer places a large leak can happen. This is the "
         "Aegis-40 layout."),
        ("SBF — Soluble-Boron-Free", None,
         "Normal PWRs dissolve boron in the coolant to soak up spare neutrons and control the "
         "reaction. Aegis-40 uses NO dissolved boron; it controls the reaction with solid burnable "
         "absorbers and control rods instead. Benefit: no boron means no boron-dilution accident and "
         "much less liquid radioactive waste. This is a defining Aegis-40 choice."),
        ("Natural circulation", None,
         "Moving the coolant using only the fact that hot water rises and cold water sinks — no "
         "pumps. Aegis-40's primary loop has NO reactor coolant pumps; the heat itself drives the "
         "flow. One less thing that can break or lose power."),
        ("Once-through fuel cycle", None,
         "Use the fuel once, then store it — no chemical reprocessing to recover usable material. "
         "Simpler and more proliferation-resistant; the trade-off is more spent fuel to store."),
        ("Passive safety", None,
         "Safety that works by physics alone (gravity, natural circulation, stored pressure) with no "
         "pumps, motors, operator action or even electricity required. The opposite of 'active' "
         "safety that needs powered equipment to run."),
     ]),

    ("2.  Power, Energy & Performance",
     "How big the plant is and how well it runs. The single most confused point: thermal vs "
     "electric power — they are NOT the same number.",
     [
        ("Thermal power [MWth]", "MWth",
         "The raw HEAT the core produces, in megawatts. Aegis-40 = 125 MWth. 'th' = thermal. This is "
         "always bigger than the electricity, because turning heat into electricity is never 100% "
         "efficient."),
        ("Electric power [MWe]", "MWe",
         "The actual ELECTRICITY sent to the grid, in megawatts. Aegis-40 = 40 MWe (net). 'e' = "
         "electric. 'Net' means after the plant's own equipment has taken its share."),
        ("Net thermal efficiency [η]", "η, %",
         "Electricity out ÷ heat in = 40 MWe ÷ 125 MWth = 32%. The Greek letter is eta (η). The other "
         "68% of the heat is rejected to the environment (via the cooling tower) — that is normal and "
         "unavoidable for a steam plant; 32% is a typical, honest number for this size."),
        ("Gross vs net", None,
         "Gross = electricity the generator makes. Net = gross minus what the plant uses to run itself "
         "(pumps, fans, lights) — the 'house load'. Only the net reaches the grid."),
        ("Capacity factor [%]", "%",
         "How much electricity the plant actually makes over a year compared with the maximum it "
         "COULD make if it ran flat-out 24/7 all year. 95% means it produces full power 95% of the "
         "time; the missing 5% is planned refuelling and maintenance outages. Higher = more revenue."),
        ("Design lifetime [years]", "yr",
         "How many years the plant is engineered to operate — here 60. Set by the one part you cannot "
         "replace: the reactor vessel, which slowly gets brittle from neutron bombardment. Low-power "
         "SMRs damage the vessel slowly, so 60 (even 80) years is realistic."),
        ("EFPD — Effective Full-Power Days", "EFPD",
         "Days of running AS IF at 100% power. A reactor that ran 200 days at half power = 100 EFPD. "
         "It is the honest way to measure how much energy the fuel has delivered. Aegis-40's cycle = "
         "479 EFPD (~16 months) between refuellings."),
        ("Burnup [GWd/tHM]", "GWd/tHM, GWd/MTU",
         "How much energy has been squeezed out of each tonne of fuel before it is discarded — "
         "gigawatt-days per tonne of heavy metal. Aegis-40 ≈ 42.8 GWd/tHM. Higher burnup = more "
         "energy per tonne of uranium = less fuel mined and less waste per unit of electricity. "
         "(GWd/MTU means the same thing: MTU = metric tonnes of uranium.)"),
        ("Specific power [MW/tHM]", "MW/tHM",
         "Heat produced per tonne of fuel = 125 MWth ÷ 5.3 t ≈ 23.6. A 'how hard is the fuel working' "
         "number; lower means a more relaxed, cooler-running core."),
        ("Heavy metal / tHM", "tHM",
         "The uranium (and any plutonium) mass in the fuel, NOT counting the oxygen or cladding. "
         "Measured in tonnes of heavy metal. Aegis-40 starts with ~5.3 tHM."),
        ("Cogeneration", None,
         "Using the reactor's heat for more than just electricity at the same time — e.g. also making "
         "hydrogen or supplying district heating. Raises the total useful energy from the same fuel."),
     ]),

    ("3.  Neutronics — the Nuclear Physics of the Core",
     "How the chain reaction is kept steady, controlled, and shut-down-able. This is the part the "
     "OpenMC computer model calculates.",
     [
        ("Neutron", None,
         "The tiny particle that carries the chain reaction. One fission releases neutrons; if on "
         "average exactly one of them goes on to cause the next fission, the reactor is steady."),
        ("Moderator", None,
         "Something that slows fast neutrons down to the gentle speeds at which they split uranium "
         "best. In a PWR the moderator is the cooling water itself — so losing water also slows the "
         "reaction, which is a built-in safety feature."),
        ("k-effective [k_eff]", "k_eff",
         "The 'multiplication factor': average number of fissions caused by each fission. k = 1.000 "
         "exactly steady (critical); k > 1 power rising (supercritical); k < 1 power falling "
         "(subcritical). The whole game is keeping k near 1 during operation and safely below 1 when "
         "shut down."),
        ("Reactivity [ρ, pcm]", "ρ, pcm",
         "How far k is from 1, as a fraction: ρ = (k−1)/k. Measured in pcm = 'per cent mille' = "
         "hundred-thousandths. 1000 pcm = 1% extra. Greek letter rho (ρ). Positive ρ pushes power up, "
         "negative pulls it down. (Note: ρ is also used for density elsewhere — context tells which.)"),
        ("Excess reactivity", None,
         "The spare reactivity built into fresh fuel so the core can run for a whole cycle as the "
         "fuel slowly depletes. It must be 'held down' at the start by absorbers (see below), or the "
         "reactor would be far too supercritical on day one."),
        ("Critical / supercritical / subcritical", None,
         "Critical = steady chain reaction (k=1), the normal running state. Supercritical = power "
         "climbing. Subcritical = power dying away — the safe state for shutdown and storage. "
         "'Critical' here is normal and good, not dangerous."),
        ("BOL / BOC / EOC", None,
         "Beginning-Of-Life, Beginning-Of-Cycle, End-Of-Cycle. Snapshots in the fuel's life: fresh "
         "(BOL/BOC) vs nearly spent and ready to refuel (EOC). Many numbers are quoted at both ends "
         "because the core behaves differently as fuel ages."),
        ("Burnable absorber (BA) — Gd, Er", None,
         "A material mixed into fresh fuel that strongly absorbs neutrons at first (holding the "
         "excess reactivity down) and then 'burns out' over time, releasing reactivity just as the "
         "fuel needs it. Aegis-40 uses gadolinium (Gd₂O₃, strong, burns out early) plus erbium "
         "(Er₂O₃, gentle, lasts longer). They replace the boron a normal PWR would dissolve in water."),
        ("Enrichment [wt%]", "wt%",
         "The percentage of the fissile isotope uranium-235 in the uranium (natural uranium is only "
         "0.7%). Aegis-40 uses 4.40–4.95 wt%, just under the 5.0% legal limit for commercial fuel. "
         "Higher enrichment lets the fuel run longer before refuelling."),
        ("Depletion / burnup calculation", None,
         "The simulation of how the fuel composition (and k) changes as it is used up over the cycle "
         "— uranium fissions away, absorbers burn out, fission products build up. Produces the "
         "k-vs-burnup curve (FER Fig 8.2-5)."),
        ("Reactivity coefficients (MTC, DTC, void)", "pcm/K, pcm/%void",
         "How reactivity changes when conditions change. Moderator Temperature Coefficient (water "
         "hotter), Doppler/fuel Temperature Coefficient (fuel hotter), Void coefficient (water turns "
         "to steam). For safety ALL must be NEGATIVE — meaning any unwanted heat-up automatically "
         "reduces power. Aegis-40's are all negative (self-stabilising)."),
        ("Doppler effect", None,
         "As fuel heats up, uranium absorbs more neutrons uselessly, which lowers power. An instant, "
         "automatic brake on power surges — one of the most important inherent safety effects."),
        ("Xenon", None,
         "Xenon-135, a fission product that is a very strong neutron absorber. It builds up after "
         "start-up and temporarily pushes reactivity down (the early 'dip' in the k curve), then "
         "settles out. A normal, expected behaviour the operators plan around."),
        ("Shutdown margin [%Δk/k]", "%Δk/k",
         "How firmly subcritical (how far below k=1) the reactor is when the control rods are dropped "
         "in — even assuming the single most effective rod is stuck out. Must be safely positive. "
         "Aegis-40 ≈ 12% — a large, comfortable margin."),
        ("Control-rod worth [pcm]", "pcm",
         "How much total reactivity the control rods can remove when fully inserted. Bigger worth = "
         "stronger ability to shut the reactor down. Aegis-40 ≈ 15,000 pcm vs a 5,000 pcm requirement."),
        ("Peaking factors (F_q, F_ΔH, F_z)", "F_q, F_ΔH, F_z",
         "Power is not perfectly even across the core; peaking factors measure the hottest spot vs "
         "the average. F_q = hottest single point in 3D (here 3.48). F_ΔH = hottest fuel channel, "
         "radially (2.27). F_z = how uneven the power is top-to-bottom. Lower/flatter is better for "
         "cooling margin. Flattening these is why the fuel is arranged in zones."),
        ("Zoning", None,
         "Deliberately arranging fuel and absorbers so the centre isn't far hotter than the edge — "
         "e.g. more gadolinium in the central assemblies. Flattens the power shape and improves "
         "cooling margin without changing total power."),
     ]),

    ("4.  Fuel & Materials",
     "What the fuel is actually made of and the limits that keep it intact.",
     [
        ("UO₂ — uranium dioxide", None,
         "The fuel: a hard black ceramic of uranium and oxygen, pressed into small pellets. Chosen "
         "because it has a very high melting point (~2840 °C) and a huge track record."),
        ("Fuel pellet", None,
         "A cylinder of UO₂ about the size of a fingertip. Thousands of them stacked end-to-end fill "
         "the fuel rods."),
        ("Fuel rod / fuel pin", None,
         "A thin sealed metal tube (~9.5 mm across, ~2 m long) packed with the pellet stack. "
         "'Rod' and 'pin' mean the same thing."),
        ("Cladding — Zircaloy (Zr-4)", None,
         "The metal tube around the pellets — a zirconium alloy. It is the FIRST barrier holding "
         "radioactive material inside the fuel, so keeping it intact is a core safety goal. Zircaloy "
         "is used because it barely absorbs neutrons and resists hot-water corrosion. (Note: PWRs use "
         "Zr-4; Zircaloy-2 is for boiling reactors — a common mix-up.)"),
        ("Fuel assembly (FA)", None,
         "A bundle of fuel rods held in a square grid — here a 17×17 array (289 positions). Aegis-40's "
         "core is 21 such assemblies. The assembly is the unit that gets loaded and removed."),
        ("Guide tube / instrument tube", None,
         "The ~25 positions in each 17×17 assembly that hold no fuel — they make room for control "
         "rods to slide in (guide tubes) and for measuring instruments (instrument tube)."),
        ("Lattice / pitch", None,
         "The regular grid the rods sit on; 'pitch' is the centre-to-centre spacing between rods "
         "(~12.6 mm). The water in the gaps is the moderator and coolant."),
        ("Linear heat rate [kW/m]", "kW/m",
         "How much power each metre of fuel rod produces. Aegis-40 averages ~11 kW/m, peak ~39 kW/m. "
         "A key limit: too high and the fuel centre could melt. Aegis-40 runs comfortably low."),
        ("Centerline temperature [°C]", "°C",
         "The temperature at the very centre of a fuel pellet (its hottest point). Must stay well "
         "below the ~2840 °C melting point; Aegis-40 peaks ~1750 °C, with ~1090 °C to spare."),
        ("Fission gas release (FGR)", None,
         "As fuel is used, some fission products are gases that escape the pellet and build up "
         "pressure inside the sealed rod. The rod has an empty space (plenum) and a spring to absorb "
         "this — a fuel-integrity design point."),
        ("Plenum", None,
         "The empty gas-collecting space at the end of a fuel rod (also a word for an open coolant "
         "space in the vessel — context tells which)."),
        ("Reload batch / 4-batch", None,
         "Fuel is not all replaced at once. At each refuelling only a fraction (here ¼ — a '4-batch' "
         "scheme) is swapped for fresh fuel, and the rest is shuffled. Evens out the power and uses "
         "fuel more fully."),
     ]),

    ("5.  Cooling & Thermal-Hydraulics (the Primary Loop)",
     "How heat gets out of the core. 'Thermal-hydraulics' just means heat + fluid flow.",
     [
        ("Primary loop / primary coolant", None,
         "The high-pressure water that flows through the core, picks up its heat, and carries it to "
         "the steam generator — then returns. It stays inside the vessel, is mildly radioactive, and "
         "never mixes with the secondary water."),
        ("Secondary loop", None,
         "The separate, clean water/steam circuit that receives heat from the primary through the "
         "steam generator wall, boils into steam, and drives the turbine. Never radioactive — that "
         "separation is deliberate."),
        ("Coolant inventory [t]", "t, kg",
         "The total mass of primary water inside the vessel (Aegis-40 ≈ 26 tonnes). More water = more "
         "stored heat-soak, so the core heats up more slowly if cooling is lost — a safety plus."),
        ("Core inlet / outlet temperature [°C]", "°C",
         "How hot the coolant is going INTO the core (cooler, ~265 °C) and coming OUT (hotter, "
         "~305 °C). The difference (ΔT, ~40 °C) times the flow tells you the power carried."),
        ("Pressuriser", None,
         "A tank that keeps the primary water at the right high pressure (~12.8 MPa) so it stays "
         "liquid and never boils in the core. It has heaters and a steam cushion on top that absorb "
         "pressure swings. In an iPWR it sits in the top of the same vessel."),
        ("Subcooling [°C]", "°C",
         "How many degrees BELOW its boiling point the water is. The primary water is kept subcooled "
         "so it stays liquid; positive subcooling = no boiling = good."),
        ("Steam generator (SG) / OTSG", None,
         "The heat exchanger where hot primary water boils the clean secondary water into steam, "
         "without the two ever touching. OTSG = Once-Through Steam Generator: secondary water enters "
         "as liquid and leaves as dry steam in a single pass through helical coil tubes. Aegis-40's "
         "SG sits inside the reactor vessel."),
        ("Downcomer / riser", None,
         "The flow path inside the vessel: hot water goes UP the central riser out of the core; "
         "cooled water comes back DOWN the outer downcomer annulus to re-enter the core. This loop is "
         "what natural circulation drives."),
        ("Plenum (lower / upper)", None,
         "The open water spaces below and above the core that collect and distribute the coolant flow."),
        ("DNBR / MDNBR", "—",
         "Departure-from-Nucleate-Boiling Ratio. A safety margin against the coolant forming an "
         "insulating steam film on the hot rods (which would let them overheat). MDNBR = the minimum "
         "(worst-spot) value; it must stay above a safe limit. The real thermal limit on power."),
        ("Saturation temperature / pressure", None,
         "The boiling point of water at a given pressure (water boils hotter when squeezed harder). "
         "At 4.8 MPa it boils at ~261 °C. Used everywhere in cycle and cooling design."),
     ]),

    ("6.  The Power-Conversion Cycle (Heat → Electricity)",
     "The 'secondary side' — a steam engine. This is the §8.9 Rankine cycle. Imagine following one "
     "drop of water around the loop.",
     [
        ("Rankine cycle", None,
         "The standard steam power cycle: boil water → spin a turbine with the steam → condense it "
         "back to water → pump it back to the boiler → repeat. Every coal, gas and nuclear steam "
         "plant uses it. Aegis-40's secondary side is a Rankine cycle."),
        ("Turbine (HP / LP)", None,
         "A fan-like machine the steam spins as it expands, turning heat energy into rotation. HP = "
         "High-Pressure (first, takes fresh steam), LP = Low-Pressure (last, takes the now-weaker "
         "steam). On the diagram they are drawn as trapezoids that widen as the steam expands."),
        ("Generator", None,
         "Spun by the turbine; converts rotation into electricity. Drawn as a circle marked 'G'."),
        ("Condenser", None,
         "A big heat exchanger after the turbine that cools the spent steam back into water by "
         "dumping its leftover heat to the cooling-water circuit. This is where most of the 'wasted' "
         "68% of heat leaves the plant."),
        ("Cooling tower", None,
         "The tall waisted (hyperboloid) structure that releases the condenser's waste heat to the "
         "air; the 'steam' you see is just water vapour. Drawn as the hourglass shape on the diagram."),
        ("Feedwater", None,
         "The condensed water on its way back to the steam generator to be boiled again. 'Feed' = "
         "feeding the boiler."),
        ("Feedwater pump", None,
         "Pushes the feedwater back up to the steam generator's high pressure. Drawn as a circle with "
         "a triangle. (These secondary-side pumps are normal — only the PRIMARY loop is pump-free.)"),
        ("Feedwater heater (FWH) / regeneration", None,
         "Heat exchangers that pre-warm the feedwater using a little steam bled from the turbine, so "
         "less fuel heat is needed to boil it. This 'regeneration' is why real cycle efficiency beats "
         "the naive number. Aegis-40 has two."),
        ("Deaerator", None,
         "A special open feedwater heater that also strips dissolved air/oxygen out of the water "
         "(oxygen causes corrosion). Drawn in the feedwater train."),
        ("Extraction / bleed steam", None,
         "Steam deliberately tapped from a mid-point of the turbine to feed the feedwater heaters. "
         "You give up a little electricity to save more fuel heat — a net win. Shown as dashed steam "
         "lines on the diagram (e.g. 'extr. 0.15 MPa')."),
        ("Moisture separator", None,
         "A device between turbine stages that removes water droplets from the steam. Wet steam "
         "erodes turbine blades, so the steam is dried before the last stage. Marked 'MS' on the PFD."),
        ("Steam quality [x]", "x",
         "The fraction of a steam/water mix that is actually steam. x = 1.0 is fully dry steam; "
         "x = 0.9 means 10% is water droplets. Kept high (≥~0.88) at the turbine exit to protect the "
         "blades."),
        ("Enthalpy [h, kJ/kg]", "h, kJ/kg",
         "The heat-energy content of the steam/water per kilogram. Differences in enthalpy across a "
         "component tell you how much energy it added or extracted. The main bookkeeping quantity of "
         "the cycle."),
        ("Entropy [s, kJ/kg·K]", "s",
         "A measure of disorder/quality of energy. Used mostly to draw the T–s diagram and to judge "
         "how close a turbine is to ideal. You don't need to feel it intuitively — just know a "
         "perfect (lossless) expansion keeps entropy constant."),
        ("T–s diagram", None,
         "Temperature-vs-entropy plot — the standard 'map' of a thermodynamic cycle. The loop's area "
         "represents the work produced. FER Fig 8.9-2."),
        ("Isentropic efficiency", None,
         "How close a real turbine or pump gets to the ideal, lossless version (e.g. 85%). Real "
         "machines always fall short; this number captures by how much."),
        ("OTSG pinch / approach", None,
         "In the steam generator, the closest temperature gap between the hot primary water and the "
         "boiling secondary water (the 'pinch'). It must stay positive and not too small, or the heat "
         "won't transfer. A feasibility check on the boiler design."),
        ("TES — Thermal Energy Storage", None,
         "A heat 'battery'. When electricity is cheap/unneeded, store the reactor's heat in big tanks "
         "of hot fluid; release it later for extra power or heat. Lets a steady reactor follow a "
         "changing grid. Aegis-40 uses a two-tank design."),
        ("Two-tank sensible-heat storage", None,
         "TES done with a cold tank and a hot tank of thermal fluid: to charge, pump fluid cold→hot "
         "through a heater; to discharge, send it hot→cold through a boiler. 'Sensible heat' = stored "
         "simply by the fluid getting hotter (no boiling/melting)."),
        ("Therminol-66 (thermal fluid)", None,
         "A heat-transfer OIL used as the storage fluid — it carries heat at high temperature without "
         "needing high pressure (unlike water). NOTE: it is a 'thermal' (heat-carrying) fluid; the "
         "early diagram label 'Therminol' was corrected to read 'thermal fluid'."),
        ("SOE — Solid-Oxide Electrolyser", None,
         "A unit that splits water into hydrogen and oxygen using electricity AND high-temperature "
         "heat. Running it on the reactor's off-peak power and steam makes hydrogen to sell — a second "
         "product. ('Electrolysis' = splitting with electricity.)"),
        ("District heating", None,
         "Piping the plant's low-grade waste heat to nearby buildings/industry for space heating — "
         "useful energy that would otherwise be dumped to the air."),
     ]),

    ("8.  Safety Systems & Risk Numbers",
     "The systems that prevent and mitigate accidents, and the numbers that score how safe the "
     "design is. Many of these are owned by the systems/safety sections.",
     [
        ("Defense in depth", None,
         "Safety by multiple independent layers, so no single failure is dangerous: keep the fuel "
         "intact, keep the coolant in, keep radioactivity in the containment — each backs up the next."),
        ("Redundancy / trains", None,
         "Having several independent copies of a safety system ('trains') so it still works if one "
         "fails. '3 redundant trains' = three full sets, any of which can do the job."),
        ("SCRAM / reactor trip", None,
         "An emergency shutdown — slam all control rods into the core to stop the chain reaction fast."),
        ("CRDM — Control-Rod Drive Mechanism", None,
         "The motor/latch assembly that positions the control rods and lets them drop on a trip. In "
         "an iPWR these sit inside the vessel."),
        ("ECCS — Emergency Core Cooling System", None,
         "Backup systems that flood the core with water if normal cooling is lost, to stop the fuel "
         "overheating. Aegis-40's are passive (work without power)."),
        ("RHR — Residual Heat Removal", None,
         "Even after shutdown the core keeps making 'decay heat' for a long time; the RHR system "
         "carries that heat away so the fuel stays cool."),
        ("Containment", None,
         "The leak-tight structure around the reactor that traps any released radioactivity as the "
         "final barrier. 'Dry' containment = filled with air/nitrogen (not water). Rated for a peak "
         "accident pressure (e.g. 4.14 bar)."),
        ("Decay heat", None,
         "Heat the fuel keeps producing after shutdown from radioactive decay of fission products — "
         "large at first (~several % of full power), fading over hours and days. It is why a reactor "
         "must be cooled even when 'off', and it drives the ECCS/RHR and spent-fuel-pool design."),
        ("EPZ — Emergency Planning Zone [km]", "km",
         "The radius around the plant needing emergency plans (evacuation, etc.). Big plants need "
         "~16 km; Aegis-40 aims for ~0.5 km (the site boundary) because its small inventory and "
         "passive safety keep any release tiny. A major SMR selling point — but it must be EARNED by "
         "the dose analysis."),
        ("SSE — Safe-Shutdown Earthquake [g]", "g",
         "The strongest earthquake the plant is designed to safely shut down through, measured in g "
         "(fraction of gravity's acceleration). 0.3 g is a robust generic value covering most sites."),
        ("CDF — Core Damage Frequency", "per reactor-year",
         "The estimated probability per year that the core is seriously damaged — i.e. how often a "
         "bad accident is expected. Aegis-40 targets < 1×10⁻⁷ (less than one in ten million years). "
         "Smaller = safer."),
        ("LRF — Large Release Frequency", "per reactor-year",
         "The estimated yearly probability of a LARGE radioactive release to the environment — the "
         "outcome people actually care about. Target < 1×10⁻⁸. Even smaller than CDF because "
         "containment catches most damage cases."),
        ("PSA / PRA", None,
         "Probabilistic Safety (or Risk) Assessment — the systematic study that produces CDF and LRF "
         "by adding up all the accident scenarios and their odds."),
     ]),

    ("9.  Waste & the Back-End",
     "What happens to the fuel after it leaves the core (FER §8.11).",
     [
        ("Spent fuel", None,
         "Fuel removed from the core after use — still highly radioactive and heat-producing, so it "
         "needs cooling and shielding."),
        ("HLW — High-Level Waste", None,
         "The most radioactive, heat-generating waste category — spent fuel is HLW. Needs robust "
         "long-term isolation."),
        ("Source term", None,
         "The full inventory of radioactive material present (how much of each isotope, its activity "
         "and heat) — the starting point for any dose, decay-heat or storage calculation."),
        ("Activity [Bq]", "Bq",
         "How many radioactive decays happen per second, in becquerels. A raw measure of 'how "
         "radioactive' something is."),
        ("Radiotoxicity [Sv]", "Sv",
         "The biological harm potential if the material got into a body, in sieverts — it weights "
         "activity by how damaging each isotope is to humans. Drops by orders of magnitude as the "
         "waste decays over decades."),
        ("Waste intensity [tHM/TWhe]", "tHM/TWhe",
         "Tonnes of spent fuel produced per unit of electricity generated. Lower = cleaner. Aegis-40 "
         "≈ 3.0, about half of the CAREM-25 reference (~6.4), thanks to high burnup and efficiency."),
        ("Cooling time", None,
         "How long spent fuel has sat since discharge. Used to schedule moving it from the wet pool "
         "to dry storage once its heat and dose have fallen enough."),
        ("Dry-cask storage", None,
         "Storing cooled spent fuel in sealed, air-cooled concrete/steel casks on-site — passive, no "
         "pumps. Follows the pool once the fuel is cool enough."),
        ("Spent-fuel pool", None,
         "A deep water pool that cools and shields freshly discharged fuel for the first years; water "
         "both removes decay heat and blocks radiation."),
        ("Storage criticality / burnup credit", None,
         "Making sure stored fuel can NEVER restart a chain reaction (stays subcritical). 'Burnup "
         "credit' = taking credit for the fact that used fuel is less reactive than fresh, which (with "
         "neutron-absorbing racks) keeps storage safely subcritical."),
        ("k(95/95)", "—",
         "The storage k_eff reported with conservative statistics: 95% confidence that 95% of cases "
         "stay below it. Must be ≤ 0.95. Aegis-40 ≈ 0.78 — a wide safety margin."),
        ("Repository", None,
         "A permanent deep-underground disposal site for HLW. Aegis-40's tiny arisings let it defer "
         "this for the plant's life via on-site storage."),
     ]),
]

# --- the PFD symbol key (rendered as a table) -------------------------------
PFD_TITLE = "7.  Reading the Plant Flow Diagram (PFD) — Symbols & Colours"
PFD_INTRO = (
    "A Process Flow Diagram (PFD) is the plant's 'circuit diagram': boxes and shapes are equipment, "
    "lines are pipes, and each pipe's colour says what is flowing in it. Below is the key for the "
    "Aegis-40 §8.9 diagram (Fig 8.9-1). Equipment also carries a short TAG (e.g. E-1, P-C) — a "
    "label so text can refer to it; the letter hints at the type (E = heat Exchanger, P = Pump, "
    "T = Turbine, R = Reactor, TK = Tank, CT = Cooling Tower).")
PFD_SYMBOLS = [
    ("Trapezoid (widening)", "Turbine (HP / LP) — steam expands and spins it; widens in the flow direction."),
    ("Circle with 'G'", "Electrical generator."),
    ("Circle with a triangle", "Pump — pushes liquid; the triangle points the flow way."),
    ("Bow-tie / two triangles", "Valve — opens, closes or throttles a flow (e.g. TCV = turbine control valve)."),
    ("Rectangle with internal lines", "Heat exchanger (condenser, feedwater heater, boiler) — the lines hint at tubes."),
    ("Tall waisted hourglass", "Cooling tower — rejects waste heat to the air."),
    ("Cylinder / tank shape", "Storage tank (e.g. the hot and cold TES tanks)."),
    ("Rounded box", "A vessel or major component (e.g. the reactor RPV, the SOE electrolyser)."),
    ("Small solid dot on a line", "A 'tee' — a pipe branch / junction where flow splits or joins."),
    ("Arrowhead on a line", "Direction the fluid flows."),
    ("Dashed line", "A secondary / tapped / minor stream (e.g. extraction steam, a drain)."),
]
PFD_COLORS = [
    ("Dark red", "Primary coolant — the radioactive water inside the reactor vessel."),
    ("Orange", "Main steam — high-energy steam going to the turbine."),
    ("Blue", "Feedwater / condensate — the liquid water returning to be boiled."),
    ("Light blue", "Cooling water — the circuit to the cooling tower (rejected heat)."),
    ("Dark red / teal (TES)", "Hot / cold THERMAL FLUID (Therminol-66) in the storage loop."),
    ("Green", "Electricity — generator output to the grid."),
    ("Purple", "Hydrogen — product from the SOE electrolyser."),
    ("Amber/brown", "District-heating supply."),
]

UNITS = [
    ("MWth", "megawatt thermal", "Heat power (core output). 125 for Aegis-40."),
    ("MWe", "megawatt electric", "Electric power to grid. 40 for Aegis-40."),
    ("MPa / bar / psi", "pressure", "1 MPa ≈ 10 bar ≈ 145 psi. Primary ~12.8 MPa; secondary ~4.8 MPa."),
    ("°C / K", "temperature", "Kelvin = °C + 273. A temperature DIFFERENCE in K and °C is the same size."),
    ("GWd/tHM", "gigawatt-days per tonne", "Fuel burnup (energy squeezed per tonne). ~42.8."),
    ("EFPD", "effective full-power days", "Days of energy as if at 100% power. Cycle ≈ 479."),
    ("pcm", "per cent mille", "Reactivity unit = 0.001%. 1000 pcm = 1%."),
    ("wt%", "weight percent", "Fraction by mass — e.g. U-235 enrichment, 4.95 wt%."),
    ("kW/m", "kilowatt per metre", "Linear heat rate along a fuel rod."),
    ("kg/s", "kilograms per second", "Mass flow rate. Primary ~580 kg/s."),
    ("kJ/kg", "kilojoule per kilogram", "Specific enthalpy (energy content per kg)."),
    ("Bq", "becquerel", "Radioactive decays per second (activity)."),
    ("Sv", "sievert", "Radiation dose / biological harm (radiotoxicity)."),
    ("g", "g-force", "Earthquake acceleration as a fraction of gravity. SSE 0.3 g."),
    ("η, ρ, Δ, σ, x", "Greek/symbols", "η efficiency · ρ reactivity OR density · Δ a change/difference · σ statistical uncertainty · x steam quality."),
]

ACRONYMS = [
    ("BA", "Burnable Absorber"), ("BOC/BOL/EOC", "Beginning-Of-Cycle / -Life, End-Of-Cycle"),
    ("BOP", "Balance Of Plant (everything outside the reactor island)"),
    ("CDF", "Core Damage Frequency"), ("CRDM", "Control-Rod Drive Mechanism"),
    ("DNBR / MDNBR", "(Minimum) Departure from Nucleate Boiling Ratio"),
    ("DTC", "Doppler (fuel) Temperature Coefficient"), ("ECCS", "Emergency Core Cooling System"),
    ("EFPD", "Effective Full-Power Days"), ("EPZ", "Emergency Planning Zone"),
    ("FA", "Fuel Assembly"), ("FER", "Final Evaluation Report"), ("FGR", "Fission Gas Release"),
    ("FWH", "Feedwater Heater"), ("HLW", "High-Level Waste"), ("HP / LP", "High- / Low-Pressure"),
    ("iPWR", "integral Pressurised Water Reactor"), ("ISI", "In-Service Inspection"),
    ("LRF", "Large Release Frequency"), ("MTC", "Moderator Temperature Coefficient"),
    ("MWe / MWth", "Megawatt electric / thermal"), ("OTSG", "Once-Through Steam Generator"),
    ("PFD", "Process Flow Diagram"), ("PSA / PRA", "Probabilistic Safety / Risk Assessment"),
    ("PWR", "Pressurised Water Reactor"), ("RHR", "Residual Heat Removal"),
    ("RPV", "Reactor Pressure Vessel"), ("SBF", "Soluble-Boron-Free"),
    ("SG", "Steam Generator"), ("SMR", "Small Modular Reactor"),
    ("SOE", "Solid-Oxide Electrolyser"), ("SSE", "Safe-Shutdown Earthquake"),
    ("TES", "Thermal Energy Storage"), ("tHM", "tonnes Heavy Metal"),
    ("UO₂", "Uranium dioxide (the fuel)"), ("Zr-4", "Zircaloy-4 (cladding alloy)"),
]

AEGIS_NUMBERS = [
    ("Thermal / electric power", "125 MWth / 40 MWe"),
    ("Net efficiency", "32%"),
    ("Reactor type", "integral PWR, soluble-boron-free, natural circulation"),
    ("Fuel assemblies", "21 (17×17 lattice)"),
    ("Enrichment (avg / max)", "4.54 / 4.95 wt% U-235"),
    ("Burnable absorbers", "Gd₂O₃ 8 wt% + Er₂O₃ 0.5 wt%"),
    ("Cladding", "Zircaloy-4"),
    ("Discharge burnup", "42.8 GWd/tHM"),
    ("Cycle length", "479 EFPD (~16 months), 4-batch"),
    ("Primary pressure / temps", "~12.8 MPa, 265 → 305 °C"),
    ("Primary flow / inventory", "~580 kg/s / ~26 t"),
    ("Secondary (steam) pressure", "4.8 MPa"),
    ("Design life / capacity factor", "60 yr / 95%"),
    ("Safety targets", "CDF < 1e-7, LRF < 1e-8 per yr, EPZ ~0.5 km, SSE 0.3 g"),
]

# ===========================================================================
# RENDERING
# ===========================================================================

def set_cell_shade(cell, hexcolor):
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), hexcolor)
    cell._tc.get_or_add_tcPr().append(sh)

def style_base(doc):
    n = doc.styles["Normal"]
    n.font.name = "Arial"; n.font.size = Pt(11); n.font.color.rgb = INK
    n.paragraph_format.space_after = Pt(4); n.paragraph_format.line_spacing = 1.12
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(2.2)
        s.left_margin = s.right_margin = Cm(2.3)

def h1(doc, text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(16); p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text); r.font.name = "Arial Black"; r.font.size = Pt(15); r.font.color.rgb = ACCENT
    p.paragraph_format.keep_with_next = True
    return p

def body(doc, text, italic=False):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(text); r.italic = italic
    if italic: r.font.size = Pt(10.5); r.font.color.rgb = RGBColor(0x55,0x55,0x55)
    return p

def term_entry(doc, term, unit, definition):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    rt = p.add_run(term); rt.bold = True; rt.font.color.rgb = ACCENT
    if unit:
        ru = p.add_run("  [" + unit + "]"); ru.italic = True; ru.font.size = Pt(10); ru.font.color.rgb = RGBColor(0x80,0x80,0x80)
    p.add_run("  —  " + definition)

def two_col_table(doc, header, rows, w0=4.0):
    t = doc.add_table(rows=1, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = "Table Grid"
    for j, htext in enumerate(header):
        c = t.rows[0].cells[j]; c.text = ""
        rr = c.paragraphs[0].add_run(htext); rr.bold = True; rr.font.size = Pt(10); rr.font.color.rgb = RGBColor(0xFF,0xFF,0xFF)
        set_cell_shade(c, "1F4E79")
    for a, b in rows:
        cells = t.add_row().cells
        r0 = cells[0].paragraphs[0].add_run(a); r0.bold = True; r0.font.size = Pt(10)
        cells[1].paragraphs[0].add_run(b).font.size = Pt(10)
    for row in t.rows:
        row.cells[0].width = Cm(w0); row.cells[1].width = Cm(11.5 - (w0-4.0))
    return t

def three_col_table(doc, header, rows):
    t = doc.add_table(rows=1, cols=3); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = "Table Grid"
    for j, htext in enumerate(header):
        c = t.rows[0].cells[j]; c.text = ""
        rr = c.paragraphs[0].add_run(htext); rr.bold = True; rr.font.size = Pt(10); rr.font.color.rgb = RGBColor(0xFF,0xFF,0xFF)
        set_cell_shade(c, "1F4E79")
    for a, b, cc in rows:
        cells = t.add_row().cells
        r0 = cells[0].paragraphs[0].add_run(a); r0.bold = True; r0.font.size = Pt(9.5)
        cells[1].paragraphs[0].add_run(b).font.size = Pt(9.5)
        cells[2].paragraphs[0].add_run(cc).font.size = Pt(9.5)
    return t

def build():
    doc = Document(); style_base(doc)

    # title
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Aegis-40 iPWR"); r.font.name = "Arial Black"; r.font.size = Pt(24); r.font.color.rgb = ACCENT
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("FER Plain-Language Handbook & Glossary"); r.font.name = "Arial Black"; r.font.size = Pt(15)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Every term, symbol and unit in the Final Evaluation Report — explained simply")
    r.italic = True; r.font.size = Pt(11); r.font.color.rgb = RGBColor(0x66,0x66,0x66)
    doc.add_paragraph()
    body(doc, INTRO)
    doc.add_page_break()

    pos = 0
    for title, intro, entries in SECTIONS:
        # insert the PFD chapter (no.7) in numerical order, between sections 6 and 8
        if title.startswith("8.") and pos == 0:
            render_pfd(doc); pos = 1
        h1(doc, title); body(doc, intro, italic=True)
        for term, unit, definition in entries:
            term_entry(doc, term, unit, definition)

    # units
    h1(doc, "10.  Units, Symbols & Greek Letters")
    body(doc, "Quick reference for the units and symbols sprinkled through the FER.", italic=True)
    three_col_table(doc, ["Unit / symbol", "Says", "Plain meaning (and Aegis-40 value)"], UNITS)

    # acronyms
    h1(doc, "11.  Acronym Dictionary (A–Z)")
    two_col_table(doc, ["Acronym", "Stands for"], ACRONYMS, w0=4.2)

    # numbers at a glance
    h1(doc, "12.  Aegis-40 Numbers at a Glance")
    body(doc, "The headline design values, all in one place. (Some are still being confirmed — see "
              "the §8.1 review notes.)", italic=True)
    two_col_table(doc, ["Parameter", "Value"], AEGIS_NUMBERS, w0=6.0)

    doc.save(OUTDOC)
    nt = len(doc.tables); np_ = len(doc.paragraphs)
    print("wrote", OUTDOC)
    print(f"  sections={len(SECTIONS)+1}  glossary tables={nt}  paragraphs={np_}")

def render_pfd(doc):
    h1(doc, PFD_TITLE); body(doc, PFD_INTRO, italic=True)
    sp = doc.add_paragraph(); r = sp.add_run("Shapes (equipment & lines)"); r.bold = True; r.font.color.rgb = ACCENT
    two_col_table(doc, ["On the diagram you see…", "…which means"], PFD_SYMBOLS, w0=5.0)
    doc.add_paragraph()
    sp = doc.add_paragraph(); r = sp.add_run("Line colours (what is flowing)"); r.bold = True; r.font.color.rgb = ACCENT
    two_col_table(doc, ["Colour", "What flows in that pipe"], PFD_COLORS, w0=4.2)

if __name__ == "__main__":
    build()
