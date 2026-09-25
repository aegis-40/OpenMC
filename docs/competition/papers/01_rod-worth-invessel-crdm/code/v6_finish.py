# -*- coding: utf-8 -*-
"""Finish the Annals version: merge the two AI declarations, add Highlights."""
import copy
import re

import docx

P = (r"D:\projects\teknofest-2026-aegis-40-ipwr\docs\competition\papers"
     r"\01_rod-worth-invessel-crdm\manuscript6_annals.docx")

# One declaration covering both rounds of assistance: the v5 drafting and
# supplementary-case scripts, and the v6 model revision, analysis and plotting.
MERGED = (
    "During the preparation of this work the authors used Claude (Anthropic) to "
    "draft and edit the manuscript text, to review the manuscript, and to write "
    "and run the model-revision, analysis and plotting scripts behind the "
    "calculations reported here. The authors defined the reference design and the "
    "cases to be analysed, verified the numerical results, and are responsible for "
    "their interpretation and for the conclusions drawn. After using this tool the "
    "authors reviewed and edited the content as needed and take full responsibility "
    "for the content of the published article.")

HIGHLIGHTS = [
    "No standard-rodlet bank meets the 1 % hot zero-power stuck-rod shutdown margin",
    "The most reactive stuck cluster is an outer one, not the highest-worth one",
    "At hot zero power an outer cluster ejected from the critical bank is worth 2.46 $",
    "At 294 K the core stays supercritical with all 16 clusters inserted",
    "Cold shutdown by boron alone needs 2,220 ppm against 3,000 ppm credited",
]

d = docx.Document(P)
paras = d.paragraphs

# ---- 1. replace the AI declaration body
i = next(n for n, p in enumerate(paras)
         if p.text.strip().startswith("Declaration of generative AI"))
body = paras[i + 1]
for r in body.runs[1:]:
    r.text = ""
body.runs[0].text = MERGED
print("AI declaration merged (%d words)" % len(MERGED.split()))

# ---- 2. insert Highlights before the Abstract heading
ab = next(n for n, p in enumerate(paras) if p.text.strip() == "Abstract")
anchor = paras[ab]


def insert_before(ref_par, text, style_src):
    new = copy.deepcopy(style_src._p)
    ref_par._p.addprevious(new)
    np = docx.text.paragraph.Paragraph(new, ref_par._parent)
    for r in np.runs[1:]:
        r.text = ""
    if np.runs:
        np.runs[0].text = text
    return np


hd = insert_before(anchor, "Highlights", anchor)
hd.style = anchor.style
sample = paras[ab + 1]                      # the abstract body, for body styling
for h in HIGHLIGHTS:
    p = insert_before(anchor, "\u2022  " + h, sample)
    p.style = sample.style

over = [h for h in HIGHLIGHTS if len(h) > 85]
print("Highlights inserted: %d  (over 85 chars: %d)" % (len(HIGHLIGHTS), len(over)))
for h in over:
    print("   OVER:", len(h), h)

d.save(P)

# ---- verify
d2 = docx.Document(P)
ps = [p.text.strip() for p in d2.paragraphs if p.text.strip()]
txt = "\n".join(ps)
print("\nleftover numeric citations:", re.findall(r"\[\d+[,\u2013\-\d]*\]", txt)[:6] or "none")
j = ps.index("Highlights")
for t in ps[j:j + 7]:
    print("   ", t[:88])
