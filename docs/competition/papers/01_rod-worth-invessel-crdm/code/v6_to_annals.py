# -*- coding: utf-8 -*-
"""Convert manuscript6_clean.docx from Elsevier numbered style to Elsevier
Harvard (author-date), as Annals of Nuclear Energy requires.

    python v6_to_annals.py            # dry run: print everything, change nothing
    python v6_to_annals.py --apply    # write manuscript6_annals.docx

Citations never span runs in this document (checked: 0 of 106 in the body,
0 in tables), so in-text replacement is a per-run text substitution.
"""
import io
import re
import sys

import docx

BASE = (r"D:\projects\teknofest-2026-aegis-40-ipwr\docs\competition\papers"
        r"\01_rod-worth-invessel-crdm")
SRC = BASE + r"\manuscript6_clean.docx"
DST = BASE + r"\manuscript6_annals.docx"
APPLY = "--apply" in sys.argv

YEAR_TODO = "[YEAR TO CONFIRM]"

# entries whose surname the generic parser cannot take, and the undated ones
OVERRIDE_KEY = {
    7:  ("Boado Magan et al.", "2009"),
    21: ("U.S. NRC", YEAR_TODO),      # 10 CFR 50 App. A - CFR edition year needed
    33: ("U.S. NRC", YEAR_TODO),      # 10 CFR 50.68     - CFR edition year needed
    34: ("OECD/NEA", YEAR_TODO),      # ICSBEP handbook  - edition year needed
}

CORP = [("International Atomic Energy Agency", "IAEA"),
        ("U.S. Nuclear Regulatory Commission", "U.S. NRC"),
        ("OECD Nuclear Energy Agency", "OECD/NEA")]

AUTH = re.compile(
    r"(?:[A-ZÀ-Ý]\.(?:-[A-ZÀ-Ý]\.)*\s*)+"
    r"((?:van der |van |de |von )?[A-ZÀ-Ý][\w'\u2019\-\u00c0-\u00ff]+"
    r"(?: [A-ZÀ-Ý][\w'\u2019\-\u00c0-\u00ff]+)?)")


def split_authors(body):
    """Return (list of (surname, initials), rest-of-entry)."""
    for name, short in CORP:
        if body.startswith(name):
            return [(short, "")], body[len(name):].lstrip(", ").strip()
    out, pos = [], 0
    for part in body.split(","):
        p = part.strip()
        m = AUTH.fullmatch(p)
        if m:
            out.append((m.group(1), p[:m.start(1)].strip()))
            pos += len(part) + 1
            continue
        if p in ("et al.", "et al"):
            out.append(("et al.", ""))
            pos += len(part) + 1
        break
    return out, body[pos:].strip()


def harvard_authors(pairs):
    if len(pairs) == 1 and pairs[0][1] == "":
        return pairs[0][0]
    bits = []
    for sur, ini in pairs:
        bits.append("et al." if sur == "et al." else
                    ("%s, %s" % (sur, ini) if ini else sur))
    return ", ".join(bits)


def intext_key(pairs, corp):
    if corp:
        return pairs[0][0]
    sur = [s for s, _ in pairs]
    if "et al." in sur or len(sur) > 2:
        return "%s et al." % sur[0]
    if len(sur) == 2:
        return "%s and %s" % (sur[0], sur[1])
    return sur[0]


def reformat_tail(tail, year):
    """Drop the '(year)' and put a period between title and source where the
    journal-volume-pages pattern is unambiguous. Leaves reports untouched."""
    t = re.sub(r"\s*\(%s\)\s*" % re.escape(year), " ", tail).strip()
    m = re.search(r",\s*([A-Z][A-Za-z.&\s]*?)\s+(\d+)\s+([\d\u2013\-]+)\.", t)
    if m:
        t = t[:m.start()] + ". " + m.group(1) + " " + m.group(2) + ", " + m.group(3) + "." + t[m.end():]
    return t


d = docx.Document(SRC)
ps = [p.text.strip() for p in d.paragraphs if p.text.strip()]
ri = [n for n, t in enumerate(ps) if t == "References"][-1]
raw = ps[ri + 1:]

# ---------------------------------------------------------------- parse
entries = {}
for r in raw:
    m = re.match(r"^\[(\d+)\]\s*(.*)$", r, re.S)
    n, body = int(m.group(1)), m.group(2).strip()
    corp = any(body.startswith(c) for c, _ in CORP)
    pairs, tail = split_authors(body)
    yrs = re.findall(r"\((\d{4})\)", body) or re.findall(r"\b((?:19|20)\d{2})\b", body)
    year = yrs[-1] if yrs else YEAR_TODO
    key = intext_key(pairs, corp)
    if n in OVERRIDE_KEY:
        key, year = OVERRIDE_KEY[n]
        if key == "Boado Magan et al.":
            pairs = [("Boado Magan", "H.")] + pairs[1:]
            tail = body.split(",", 1)[1].strip() if "," in body else body
            pairs, tail = split_authors(body.replace("H. Boado Magan", "H. BoadoMagan"))
            pairs = [("Boado Magan", "H.")] + pairs[1:]
    entries[n] = dict(n=n, key=key, year=year, pairs=pairs, tail=tail, corp=corp, raw=body)

# --------------------------------------------------- a/b/c disambiguation
groups = {}
for e in entries.values():
    groups.setdefault((e["key"], e["year"]), []).append(e["n"])
for (key, year), ns in groups.items():
    if len(ns) < 2 or year == YEAR_TODO:
        continue
    for suffix, n in zip("abcdefgh", sorted(ns, key=lambda x: entries[x]["tail"].lower())):
        entries[n]["year"] = year + suffix

for e in entries.values():
    e["cite"] = "%s, %s" % (e["key"], e["year"])
    e["full"] = "%s, %s. %s" % (harvard_authors(e["pairs"]), e["year"],
                                reformat_tail(e["tail"], re.sub(r"[a-z]$", "", e["year"])))
    e["sort"] = (re.sub(r"^(van der |van |de |von )", "", e["pairs"][0][0]).lower(), e["year"])

# ------------------------------------------------------------ in-text
def expand(tok):
    out = []
    for part in re.split(r"\s*,\s*", tok):
        m = re.match(r"^(\d+)\s*[\u2013\-]\s*(\d+)$", part)
        if m:
            out += list(range(int(m.group(1)), int(m.group(2)) + 1))
        elif part.strip().isdigit():
            out.append(int(part))
    return out

CIT = re.compile(r"\[(\d+(?:\s*[,\u2013\-]\s*\d+)*)\]")

def sub_cite(m):
    ns = expand(m.group(1))
    if not ns or any(n not in entries for n in ns):
        return m.group(0)
    return "(" + "; ".join(entries[n]["cite"] for n in ns) + ")"

n_sub = 0
def convert_paragraphs(paras):
    global n_sub
    for p in paras:
        for r in p.runs:
            if "[" in r.text:
                new = CIT.sub(sub_cite, r.text)
                if new != r.text:
                    n_sub += len(CIT.findall(r.text))
                    if APPLY:
                        r.text = new

body_paras = [p for p in d.paragraphs[:]]
ref_start = None
seen = 0
for idx, p in enumerate(d.paragraphs):
    if p.text.strip() == "References":
        ref_start = idx
convert_paragraphs([p for i, p in enumerate(d.paragraphs) if ref_start is None or i < ref_start])
for t in d.tables:
    for row in t.rows:
        for c in row.cells:
            convert_paragraphs(c.paragraphs)

# ------------------------------------------------------- reference list
ordered = sorted(entries.values(), key=lambda e: e["sort"])
if APPLY:
    ref_paras = [p for i, p in enumerate(d.paragraphs) if i > ref_start and p.text.strip()]
    for p, e in zip(ref_paras, ordered):
        for r in p.runs[1:]:
            r.text = ""
        p.runs[0].text = e["full"]
    for p in ref_paras[len(ordered):]:
        for r in p.runs:
            r.text = ""
    d.save(DST)

# ----------------------------------------------------------------- report
print("in-text citations converted: %d" % n_sub)
print("reference entries          : %d\n" % len(ordered))
todo = [e for e in ordered if YEAR_TODO in e["year"]]
sfx = [e for e in ordered if re.search(r"\d{4}[a-z]$", e["year"])]
print("--- suffixed (same author+year) ---")
for e in sfx:
    print("   [%2d] %-34s %s" % (e["n"], e["cite"], e["tail"][:52]))
print("\n--- YEAR STILL NEEDED ---")
for e in todo:
    print("   [%2d] %s" % (e["n"], e["raw"][:95]))
print("\n--- first 6 Harvard entries ---")
for e in ordered[:6]:
    print("   " + e["full"][:150])
print("\n" + ("WROTE " + DST if APPLY else "dry run - nothing written"))
