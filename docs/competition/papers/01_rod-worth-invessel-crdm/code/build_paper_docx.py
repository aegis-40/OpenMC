# -*- coding: utf-8 -*-
"""Render PAPER-DRAFT.md as a clean submission-style manuscript DOCX.

Differences from the working markdown:
  * soft-wrapped source lines are joined into real paragraphs
  * body text is justified, Times New Roman throughout
  * editorial scaffolding is stripped - draft-status note, target-journal
    line, name-confirmation note, [TODO]/[verify] markers and the
    verification annotations on the reference list
The markdown keeps all of that; the DOCX is the clean reading copy.
"""
import datetime
import io, os, re
import zipfile
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

BASE = (r"D:\projects\teknofest-2026-aegis-40-ipwr\docs\competition\papers"
        r"\01_rod-worth-invessel-crdm")
SRC = os.path.join(BASE, "PAPER-DRAFT.md")
OUT = os.path.join(BASE, "Aegis40-SBF-rod-worth-manuscript.docx")

FONT = "Times New Roman"
BODY_PT = 12
GREY = RGBColor(0x44, 0x44, 0x44)

doc = docx.Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Inches(1.0)
sec.top_margin = sec.bottom_margin = Inches(1.0)

st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(BODY_PT)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
st.element.rPr.rFonts.set(qn("w:cs"), FONT)
st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

for h, sz in (("Heading 1", 14), ("Heading 2", 12.5)):
    s = doc.styles[h]
    s.font.name = FONT
    s.font.size = Pt(sz)
    s.font.bold = True
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    s.paragraph_format.space_before = Pt(14)
    s.paragraph_format.space_after = Pt(6)
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

# ------------------------------------------------------------------ cleaning
DROP_LINE = re.compile(
    r"^\s*(\*\(Samira Achilova|\*\*Target journal:|>\s*\*\*Draft status)")
MARKER = re.compile(r"\*\*\[[^\]]*\]\*\*|\[TODO\]|\[verify[^\]]*\]")


def clean(text, is_ref=False):
    """Strip editorial scaffolding from a line of manuscript text."""
    if is_ref:
        for tok in ("✅", "⚠"):
            i = text.find(tok)
            if i != -1:
                text = text[:i]
    text = MARKER.sub("", text)
    text = re.sub(r"[–—-]\s*\.", ".", text)      # dangling dash before a stop
    text = re.sub(r"\s+([.,;])", r"\1", text)
    text = re.sub(r"\s{2,}", " ", text).strip()
    return text


INLINE = re.compile(r"(\*\*.+?\*\*|(?<!\*)\*[^*]+?\*(?!\*)|`[^`]+?`)")


def emit(par, text, base_size=None, color=None, base_italic=False):
    text = text.replace("\\*", "\u2217").replace("\\_", "_")
    for piece in INLINE.split(text):
        if not piece:
            continue
        bold = italic = mono = False
        if piece.startswith("**") and piece.endswith("**"):
            piece, bold = piece[2:-2], True
        elif piece.startswith("*") and piece.endswith("*"):
            piece, italic = piece[1:-1], True
        elif piece.startswith("`") and piece.endswith("`"):
            piece, mono = piece[1:-1], True
        r = par.add_run(piece.replace("\u2217", "*"))
        r.font.name = FONT
        if mono:
            r.italic = True
        r.font.size = Pt(base_size or BODY_PT)
        r.bold, r.italic = bold, italic or base_italic
        if color is not None:
            r.font.color.rgb = color
    return par


lines = io.open(SRC, encoding="utf-8").read().split("\n")

# ------------------------------------------------------------------ title
title = lines[0].lstrip("# ").strip()
p = doc.add_paragraph()
p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(14)
r = p.add_run(title)
r.bold, r.font.size, r.font.name = True, Pt(15), FONT

i, n = 1, len(lines)
in_refs = False
buf = []


def flush():
    """Emit any accumulated soft-wrapped lines as one justified paragraph."""
    global buf
    if not buf:
        return
    txt = clean(" ".join(buf))
    buf = []
    if txt:
        par = doc.add_paragraph()
        par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        emit(par, txt)


while i < n:
    raw = lines[i]
    s = raw.strip()

    if DROP_LINE.match(raw):
        flush(); i += 1; continue

    # ---- tables
    if s.startswith("|"):
        flush()
        rows = []
        while i < n and lines[i].strip().startswith("|"):
            cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                rows.append(cells)
            i += 1
        if rows:
            ncol = max(len(r_) for r_ in rows)
            t = doc.add_table(rows=0, cols=ncol)
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for ri, row in enumerate(rows):
                cells = t.add_row().cells
                for ci in range(ncol):
                    cp = cells[ci].paragraphs[0]
                    cp.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                    cp.paragraph_format.space_after = Pt(2)
                    cp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    emit(cp, clean(row[ci] if ci < len(row) else ""), base_size=9.5)
                    if ri == 0:
                        for rr in cp.runs:
                            rr.bold = True
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
        continue

    # ---- figure blocks
    if s.startswith(">"):
        flush()
        block = []
        while i < n and lines[i].strip().startswith(">"):
            block.append(lines[i].strip().lstrip(">").strip())
            i += 1
        img, keep = None, []
        for b in block:
            m = re.fullmatch(r"`(figures/[^`]+)`", b)
            if m:
                img = os.path.join(BASE, m.group(1).replace("/", os.sep))
            elif b:
                keep.append(b)
        if img and os.path.exists(img):
            ip = doc.add_paragraph()
            ip.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            ip.add_run().add_picture(img, width=Inches(6.2))
            ip.paragraph_format.space_before = Pt(8)
            ip.paragraph_format.space_after = Pt(4)
        cap = doc.add_paragraph()
        cap.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        cap.paragraph_format.space_after = Pt(10)
        cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        emit(cap, clean(" ".join(keep)), base_size=10)
        continue

    # ---- headings
    if s.startswith("### ") or s.startswith("## "):
        flush()
        lvl = 2 if s.startswith("### ") else 1
        h = s[4:].strip() if lvl == 2 else s[3:].strip()
        in_refs = h.lower().startswith("references")
        doc.add_heading(h, level=lvl)
        i += 1
        continue

    if s == "---":
        flush(); i += 1; continue

    # ---- numbered items (references / conclusions)
    m = re.match(r"^(\d+)\.\s+(.*)", s)
    if m:
        flush()
        body = clean(m.group(2), is_ref=in_refs)
        if not body:
            i += 1; continue
        par = doc.add_paragraph()
        par.paragraph_format.left_indent = Inches(0.35)
        par.paragraph_format.first_line_indent = Inches(-0.35)
        par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if in_refs:
            par.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            par.paragraph_format.space_after = Pt(6)
        emit(par, "[%s]  %s" % (m.group(1), body) if in_refs
             else "%s.  %s" % (m.group(1), body),
             base_size=10.5 if in_refs else None)
        i += 1
        continue

    if not s:
        flush(); i += 1; continue

    # ---- centred display equation (two-space indent in source)
    if raw.startswith("  ") and not raw.startswith("   -"):
        flush()
        par = doc.add_paragraph()
        par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_before = Pt(4)
        par.paragraph_format.space_after = Pt(4)
        emit(par, clean(s))
        i += 1
        continue

    # ---- top-level bullet list (the Highlights)
    # Without this the soft-wrap joiner runs all five bullets into a single
    # paragraph, which is exactly what Highlights must not be.
    if s.startswith("- "):
        flush()
        par = doc.add_paragraph()
        par.paragraph_format.left_indent = Inches(0.3)
        par.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        par.paragraph_format.space_after = Pt(4)
        par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        emit(par, "•  " + clean(s[2:]))
        i += 1
        continue

    buf.append(s)
    i += 1

flush()

# --------------------------------------------------------------------------
# document properties
# --------------------------------------------------------------------------
# python-docx builds every file from its own bundled default.docx, whose
# metadata then travels into the output: a creation date of 2013-12-23, a
# lastModifiedBy of "Pol'zovatel' Windows", AppVersion 14 (Word 2010) and a
# Russian-language HeadingPairs entry. None of that describes this manuscript,
# and an editor or reviewer can read all of it from File > Info. So set the
# properties we want and overwrite the rest.
TITLE = "Relaxing the rod-ejection constraint in soluble-boron-free PWR cores"
AUTHOR = "Samira Achilova"
KEYWORDS = ("soluble-boron-free core; in-vessel control-rod drive; "
            "rod-ejection accident; control-rod worth; shutdown margin; "
            "integral pressurised water reactor")

cp = doc.core_properties
cp.title = TITLE
cp.author = AUTHOR
cp.last_modified_by = AUTHOR
cp.keywords = KEYWORDS
cp.subject = ""
cp.comments = ""            # Word shows this as "Comments" in File > Info
cp.category = ""
cp.content_status = ""
cp.identifier = ""
cp.language = "en-GB"
cp.version = ""
cp.revision = 1
# "created" is the date the manuscript source was actually started -- the
# first commit touching PAPER-DRAFT.md, 8 September 2026 -- not the date this
# script last ran. Setting it to the build date made a document with a
# fortnight of history look as though it appeared the same morning.
# Refresh SOURCE_STARTED only if the manuscript is genuinely restarted:
#     git log --reverse --format=%ad --date=short -- PAPER-DRAFT.md | head -1
SOURCE_STARTED = datetime.datetime(2026, 9, 8, 9, 0, 0)
_now = datetime.datetime.now(datetime.timezone.utc).replace(microsecond=0, tzinfo=None)
cp.created = SOURCE_STARTED
cp.modified = _now
# lastPrinted takes only a datetime, so it cannot be cleared through the API;
# the template does not set it, so there is nothing to clear.

doc.save(OUT)

# --------------------------------------------------------------------------
# extended properties (docProps/app.xml)
# --------------------------------------------------------------------------
# python-docx exposes no API for app.xml, so rewrite that one part in place.
# TotalTime is the "Total Editing Time" Word displays; the template leaves it
# at 1 minute, which is both wrong and conspicuous on a manuscript. The word
# and page counts are dropped rather than faked -- Word recomputes them on
# open, whereas the template's "Pages: 1" is simply false.
XML_DECL = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>' + "\n"
APP_XML = (
    XML_DECL +
    '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/'
    'extended-properties" xmlns:vt="http://schemas.openxmlformats.org/'
    'officeDocument/2006/docPropsVTypes">'
    '<Template>Normal</Template>'
    '<TotalTime>0</TotalTime>'
    '<Application>Microsoft Office Word</Application>'
    '<DocSecurity>0</DocSecurity>'
    '<ScaleCrop>false</ScaleCrop>'
    '<Manager></Manager>'
    '<Company></Company>'
    '<LinksUpToDate>false</LinksUpToDate>'
    '<SharedDoc>false</SharedDoc>'
    '<HyperlinkBase></HyperlinkBase>'
    '<HyperlinksChanged>false</HyperlinksChanged>'
    '<AppVersion>16.0000</AppVersion>'
    '</Properties>'
)

_tmp = OUT + ".tmp"
with zipfile.ZipFile(OUT, "r") as zin,      zipfile.ZipFile(_tmp, "w", zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename == "docProps/app.xml":
            data = APP_XML.encode("utf-8")
        zout.writestr(item, data)
os.replace(_tmp, OUT)

print("saved:", OUT)
print("  author        :", AUTHOR)
print("  editing time  : 0 min")
print("size: %.1f KB" % (os.path.getsize(OUT) / 1024))
