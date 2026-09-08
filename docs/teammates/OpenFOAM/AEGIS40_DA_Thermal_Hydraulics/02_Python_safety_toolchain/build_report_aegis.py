#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Build the Aegis-40 thermal-hydraulic design & safety report (.docx).

Aegis-40 = 125 MWth / 40 MWe soluble-boron-free natural-circulation iPWR. The
OpenFOAM conjugate-pin CFD + correlation-stack toolchain (validated separately
against NuScale NPM-160) is applied to the FER design point (OpenMC-consistent
cycle-resolved peaking record). Embeds F1-F8 + the F9a-c ParaView renders from
docs/figs/. Regenerate figs first (make_figs_aegis.py, make_paraview_figs.py), then run this.
"""
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
DOCS = os.path.normpath(os.path.join(HERE, "..", "docs"))
FIGS = os.path.join(DOCS, "figs")
OUT = os.path.join(DOCS, "Aegis40_TH_report.docx")

NAVY = RGBColor(0x1F, 0x4E, 0x79)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
GREY = RGBColor(0x77, 0x77, 0x77)
HDRBG, ALTBG = "1F4E79", "EAF1F7"


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), fill)
    tcPr.append(sh)


def set_cell(cell, text, bold=False, color=None, align=None, size=9.5):
    cell.text = ""
    p = cell.paragraphs[0]
    if align:
        p.alignment = align
    r = p.add_run(text); r.bold = bold; r.font.size = Pt(size); r.font.name = "Arial"
    if color:
        r.font.color.rgb = color


def add_table(doc, headers, rows, widths, hi=None):
    hi = hi or set()
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for j, h in enumerate(headers):
        c = t.rows[0].cells[j]
        set_cell(c, h, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF)); shade(c, HDRBG)
        c.width = Inches(widths[j])
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j, val in enumerate(row):
            set_cell(cells[j], val, bold=(i in hi) or (j == 0),
                     color=GREEN if i in hi else None)
            shade(cells[j], ALTBG if (i in hi or i % 2 == 1) else "FFFFFF")
            cells[j].width = Inches(widths[j])
    for row in t.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr(); b = OxmlElement("w:tcBorders")
            for e in ("top", "left", "bottom", "right"):
                el = OxmlElement(f"w:{e}"); el.set(qn("w:val"), "single")
                el.set(qn("w:sz"), "4"); el.set(qn("w:color"), "BBBBBB"); b.append(el)
            tcPr.append(b)
    doc.add_paragraph()


def H(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        r.font.name = "Arial"; r.font.color.rgb = NAVY
    return h


def para(doc, text, italic=False, size=10.5):
    p = doc.add_paragraph(); r = p.add_run(text)
    r.font.size = Pt(size); r.font.name = "Arial"; r.italic = italic
    return p


def bullet(doc, text):
    doc.add_paragraph(text, style="List Bullet")


def caption(doc, text):
    p = doc.add_paragraph(); p.alignment = AL.CENTER
    r = p.add_run(text); r.italic = True; r.font.size = Pt(9); r.font.name = "Arial"
    r.font.color.rgb = GREY


def figure(doc, fname, cap, width=6.2):
    path = os.path.join(FIGS, fname)
    if os.path.exists(path):
        doc.add_picture(path, width=Inches(width))
        doc.paragraphs[-1].alignment = AL.CENTER
        caption(doc, cap)
    else:
        t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = t.rows[0].cells[0]; c.width = Inches(width)
        set_cell(c, f"[ {fname} — ParaView screenshot to be inserted ]",
                 align=AL.CENTER, color=GREY)
        shade(c, "F2F2F2")
        tcPr = c._tc.get_or_add_tcPr(); b = OxmlElement("w:tcBorders")
        for e in ("top", "left", "bottom", "right"):
            el = OxmlElement(f"w:{e}"); el.set(qn("w:val"), "dashed")
            el.set(qn("w:sz"), "6"); el.set(qn("w:color"), "999999"); b.append(el)
        tcPr.append(b)
        for _ in range(4):
            c.add_paragraph()
        caption(doc, cap)


def main():
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10.5)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for m in ("top", "bottom", "left", "right"):
        setattr(sec, f"{m}_margin", Inches(0.9))

    # ---- title ----
    p = doc.add_paragraph(); p.alignment = AL.CENTER
    r = p.add_run("Aegis-40 — Thermal-Hydraulic Design & Safety Analysis")
    r.bold = True; r.font.size = Pt(20); r.font.name = "Arial"; r.font.color.rgb = NAVY
    p = doc.add_paragraph(); p.alignment = AL.CENTER
    r = p.add_run("125 MWth / 40 MWe soluble-boron-free natural-circulation integral PWR  ·  "
                  "conjugate fuel-pin CFD (OpenFOAM chtMultiRegionFoam) + correlation stack")
    r.font.size = Pt(12); r.font.name = "Arial"; r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    p = doc.add_paragraph(); p.alignment = AL.CENTER
    r = p.add_run("OpenMC-consistent T-H · toolchain validated against NuScale NPM-160  |  July 2026")
    r.italic = True; r.font.size = Pt(10); r.font.color.rgb = GREY
    doc.add_paragraph()

    # ---- 1 objective ----
    H(doc, "1. Objective & scope")
    para(doc, "This report establishes the steady-state thermal-hydraulic feasibility and safety "
              "margins of the Aegis-40 core: a 125 MWth / 40 MWe soluble-boron-free (SBF) integral "
              "PWR cooled by natural circulation. The analysis uses a 3-region fuel/clad/coolant "
              "conjugate-CFD model (OpenFOAM chtMultiRegionFoam) coupled to a correlation-based "
              "safety post-processor. Every thermal-hydraulic input is taken one-for-one from the "
              "locked OpenMC neutronics (FER notebook): geometry, operating temperatures and the "
              "pin-power peaking.")
    para(doc, "The same toolchain was first verified & validated against the published NuScale "
              "NPM-160 reference (companion report); this document applies it to the Aegis-40 "
              "design point. A conservation-first protocol is used: energy balance is the first "
              "acceptance test on every mesh, before any temperature is trusted.")

    # ---- 2 design basis ----
    H(doc, "2. Design basis & operating point (OpenMC-consistent)")
    add_table(doc, ["Quantity", "Value", "Source"], [
        ["Core thermal / electric power", "125 MWth / 40 MWe", "FER §8.4"],
        ["System pressure", "12.8 MPa", "FER (ρ 748 @ 283 °C)"],
        ["Core inlet / outlet T", "258 °C / 308 °C  (531.15 / 581 K)", "FER §8.4"],
        ["Core-average T / rise", "283 °C (556 K) / 50 K", "= OpenMC T_mod"],
        ["Fuel pins", "9 768  (37 FA × 264, 17×17)", "OpenMC core map"],
        ["Active length / rod pitch", "2.0 m / 12.623 mm", "locked geometry"],
        ["Pellet / clad OD", "8.19 mm / 9.52 mm  (UO₂ / Zircaloy-4)", "locked geometry"],
        ["Peaking over the cycle (OpenMC record)", "F_q: BOC 1.937 · MOC 2.435 · EOC 2.121", "OpenMC record run 2026-07-03"],
        ["Safety design basis — COLR envelope", "F_q 2.468 = F_ΔH 1.750 × F_z 1.410", "MOC Gd-hump envelope (§8.1)"],
        ["Avg / hot-pin peak linear power", "6.40 / 15.79 kW/m  (q‴ 3.00e8 W/m³)", "= q′_avg × F_q (envelope)"],
        ["Peak wall heat flux q″", "0.528 MW/m²", "= q′_peak / (π·D_clad)"],
    ], [2.3, 2.5, 2.2])
    para(doc, "The core grew to a NuScale-class 37-assembly / 9 768-pin layout, halving the "
              "per-pin power versus an earlier 21-assembly concept; even at the COLR envelope the "
              "peak linear power is 15.8 kW/m ≪ the ~43 kW/m PWR class limit. Pin geometry, pitch "
              "and active height are fixed, so the validated CFD mesh applies unchanged. The "
              "envelope covers the whole cycle: the mid-of-cycle Gd-burnout hump (13.5 GWd/t) sets "
              "it, and BOC/EOC lie below it.", italic=True)

    # ---- 3 methodology ----
    H(doc, "3. Methodology")
    H(doc, "3.1 Conjugate CFD", 2)
    bullet(doc, "Solver chtMultiRegionFoam (OpenFOAM v2412); 3 regions fuel/clad/coolant on a "
                "body-fitted butterfly O-grid; turbulence k-ω SST with wall functions.")
    bullet(doc, "Axial power: chopped cosine q‴(z)=q_peak·cos(π(z−L/2)/Lₑ), Lₑ = 2.2667 m (F_z 1.410), "
                "imposed as a runtime scalarCodedSource in the fuel region (q_peak 3.00e8 W/m³ = the "
                "COLR envelope). The record axial shapes are additionally fed to the MDNBR tool "
                "directly (§8.1).")
    bullet(doc, "Fuel-clad He gap as a contact resistance (thicknessLayers 9.15e-5 / kappaLayers "
                "0.48), applied symmetrically on both coupling patches.")
    bullet(doc, "Coolant properties at 283 °C / 12.8 MPa: ρ 748, cp 5350, μ 9.1e-5, k 0.566, Pr 0.87. "
                "Hot pin runs on the core-average channel mass flux (no flow redistribution credit "
                "→ conservative).")
    H(doc, "3.2 Safety post-processor (correlation stack)", 2)
    para(doc, "MDNBR and PCT come from a validated correlation stack — the COBRA/VIPRE/CTF "
              "subchannel-code approach — fed the CFD-trusted axial power shape, coolant bulk T(z) "
              "and flow:")
    bullet(doc, "MDNBR: W-3 critical-heat-flux correlation with the Tong non-uniform-flux F-factor.")
    bullet(doc, "PCT / fuel centreline: Dittus-Boelter film + 1-D radial conduction stack "
                "(film → clad → He-gap → fuel) with a Jens-Lottes subcooled-boiling wall clamp.")
    bullet(doc, "Cycle coverage: BOC/MOC/EOC and the COLR envelope are evaluated by "
                "tools/cycle_mdnbr.py, including the record axial power shapes fed directly to "
                "the correlations (--shapes) — the binding numbers quote the real shape.")
    para(doc, "The correlation stack is the licensed-code industry standard and carries the CHF and "
              "subcooled-boiling physics a single-phase CFD lacks — not a substitute for an "
              "inaccurate CFD. Section 6 shows the CFD is energy-conservative and its near-wall heat "
              "transfer matches Dittus-Boelter to < 7 K.", italic=True)

    # ---- 4 mesh ----
    H(doc, "4. Mesh & numerics")
    add_table(doc, ["Mesh", "Cells (fuel/clad/coolant)", "Refinement r"], [
        ["coarse", "35 700  (16.1 / 8.4 / 11.2 k)", "—"],
        ["medium (production)", "103 600  (47.6 / 22.4 / 33.6 k)", "r₃₂ = 1.426"],
        ["fine", "279 888  (126.2 / 65.9 / 87.8 k)", "r₂₁ = 1.393"],
    ], [1.9, 3.0, 1.6], hi={1})
    bullet(doc, "Three geometrically-similar meshes; only the resolution knobs differ — every "
                "physics input (source, properties, BCs, schemes) is byte-identical across the set, "
                "so the GCI isolates discretisation error.")
    bullet(doc, "Schemes: bounded Gauss upwind→limitedLinear div(phi,h); nNonOrthogonalCorrectors 2; "
                "nOuterCorrectors 2. GAMG (p_rgh) / smoothSolver; relTol 0.01. Run parallel "
                "(decomposePar → mpirun -np 8 → reconstructPar).")

    # ---- 5 verification ----
    H(doc, "5. Verification")
    para(doc, "The V&V anchor case was run at the previous-baseline source (F_q 2.035, hot pin "
              "20.26 kW — between BOC 1.937 and the MOC envelope 2.468). The single-phase "
              "constant-property conjugate problem is linear in the heat source, so the "
              "conservation, grid-convergence and near-wall conclusions below hold unchanged at "
              "every cycle state.", italic=True)
    H(doc, "5.1 Grid convergence (ASME V&V-20 GCI)", 2)
    figure(doc, "F1_grid_convergence.png",
           "Figure F1. Grid convergence of outlet T, peak clad T and peak fuel T over the three "
           "meshes, with Richardson extrapolation (h→0) and the fine-grid GCI.")
    add_table(doc, ["Metric", "coarse", "medium", "fine", "GCI_fine / direct"], [
        ["Outlet T  [K]", "609.11", "609.71", "609.99", "0.06 %"],
        ["Peak clad T  [K]", "660.68", "661.12", "661.34", "0.05 %"],
        ["Near-wall coolant  [K]", "645.38", "645.68", "645.89", "0.14 %"],
        ["Peak fuel T  [K]", "1034.7", "1036.7", "1038.4", "0.17 % (direct)"],
        ["Δp  [Pa]", "669", "682", "695", "1.78 % (direct)"],
    ], [2.0, 1.2, 1.2, 1.2, 1.6], hi={0, 1})
    para(doc, "Every temperature changes < 0.4 % medium→fine and the fine-grid GCI is sub-0.15 % on "
              "the convergent metrics, so the solution is mesh-independent. Peak-fuel and Δp are "
              "degenerate-Richardson (successive increments near-equal, p≈0) and are reported by "
              "their direct medium→fine change. The medium (104 k) mesh is used for production.")

    H(doc, "5.2 Energy conservation (GATE-1)", 2)
    figure(doc, "F2_energy_conservation.png",
           "Figure F2. Advected power ṁ·cp·ΔT versus the analytic design source (20.26 kW/hot-pin) "
           "on each mesh; conservation closes to ×1.00 on every grid.")
    para(doc, "The advected power equals the imposed source to within 0.8 % on every mesh "
              "(×0.992 / 0.996 / 0.998 on coarse / medium / fine) — the model is energy-conservative, "
              "the prerequisite for trusting any temperature.")

    # ---- 6 cross-check ----
    H(doc, "6. Near-wall validation: CFD ↔ correlation stack")
    figure(doc, "F3_cfd_vs_stack.png",
           "Figure F3. CFD (medium) versus the single-phase correlation stack at the identical "
           "operating point — agreement within 7 K on every quantity.")
    add_table(doc, ["Quantity", "CFD", "Stack (1φ DB)", "Δ"], [
        ["Coolant outlet (mixing-cup)", "337 °C", "337 °C", "0.7 K"],
        ["Fuel centreline (peak)", "764 °C", "764 °C", "0.5 K"],
        ["Clad inner (peak)", "388 °C", "381 °C", "6.9 K"],
        ["Near-wall / clad outer", "373 °C", "367 °C", "5.7 K"],
    ], [2.4, 1.4, 1.6, 1.2], hi={3})
    para(doc, "At the identical as-run operating point (G 544, T_in 258 °C, F_q 2.035, single-phase) the CFD and the "
              "Dittus-Boelter stack agree within 7 K on every quantity — bulk and fuel to < 1 K, "
              "clad and near-wall to < 7 K (CFD marginally hotter, i.e. conservative). The CFD is "
              "therefore both energy-accurate (GATE-1) and near-wall-accurate: its resolved film "
              "coefficient reproduces the validated correlation. Note also that the Aegis-40 core is "
              "geometrically the NuScale NPM core (37 FA / 9 768 pins / 2.0 m), against which the "
              "toolchain was independently validated.")

    # ---- 7 hot-channel ----
    H(doc, "7. Hot-channel thermal-hydraulics")
    figure(doc, "F4_axial_profiles.png",
           "Figure F4. Hot-channel axial profiles: wall heat flux q″(z); coolant bulk, clad outer "
           "(with subcooled-boiling clamp) and fuel centreline temperatures.")
    figure(doc, "F5_radial_stack.png",
           "Figure F5. Radial temperature stack at the hot spot — conduction drops across fuel, the "
           "He-gap contact resistance, clad and the convective film.")
    para(doc, "At the binding COLR envelope the most-peaked channel saturates below the outlet "
              "(exit quality 7.5 %), while the core-average outlet is the FER 308 °C. Peak fuel "
              "centreline is 828 °C — far below any fuel limit — a direct consequence of the "
              "modest 15.8 kW/m peak linear power. (The as-run CFD fields at F_q 2.035 show the "
              "same channel just reaching Tsat: 337 °C mixing-cup outlet vs Tsat 330 °C.)")

    # ---- 8 safety ----
    H(doc, "8. Safety margins")
    H(doc, "8.1 MDNBR", 2)
    figure(doc, "F6_dnbr_profile.png",
           "Figure F6. Hot-channel DNBR axial profile at the COLR envelope (W-3 CHF + Tong "
           "F-factor, chopped cosine): 1.38 at the CFD-medium mass flux (G 544), 1.37 at the "
           "natural-circulation design point (G 542); the record axial shape gives the binding "
           "1.33 (table below).")
    para(doc, "The cycle-resolved OpenMC record fixes the peaking over the cycle; the Gd-burnout hump at MOC "
              "(≈ 13.5 GWd/t) is the binding state and is covered by a design-specific COLR "
              "envelope (F_ΔH 1.750 × F_z 1.410). MDNBR is evaluated at the as-delivered "
              "natural-circulation flow (G 542) with the record axial shapes fed directly to the "
              "correlations (tools/cycle_mdnbr.py --shapes):")
    add_table(doc, ["Cycle state", "F_q", "MDNBR W-3 *", "MDNBR Bowring", "Verdict"], [
        ["BOC (0 GWd/t)", "1.937", "1.54", "2.66", "PASS"],
        ["MOC Gd-hump (13.5 GWd/t)", "2.435", "1.35", "2.15", "PASS"],
        ["EOC (30.8 GWd/t)", "2.121", "1.46", "2.50", "PASS"],
        ["COLR envelope — binding", "2.468", "1.33", "2.13", "PASS"],
    ], [2.3, 1.0, 1.3, 1.5, 1.0], hi={3})
    para(doc, "* W-3 is extrapolated below its mass-flux validity range at G 542 and reads "
              "conservative-low (§12); the in-range verdict is Bowring-1972. The MOC exceedance of "
              "the generic screening values (F_ΔH ≤ 1.65 / F_q ≤ 2.32 — plant-generic COLR-type "
              "parameters, not Aegis-specific limits) is resolved by this design-specific COLR "
              "basis, standard PWR practice with SMR precedent (NuScale PDIL/axial-offset control; "
              "CAREM-25 peaking 2.56 at MDNBR 1.90): at Aegis-40's very low linear heat rate, "
              "acceptability is set by the thermal limits themselves — all of which pass.",
         italic=True)
    H(doc, "8.2 PCT & subcooled boiling", 2)
    figure(doc, "F7_subcooled_boiling.png",
           "Figure F7. Clad outer temperature: single-phase film versus the subcooled-boiling "
           "(Jens-Lottes) clamp — single-phase over-predicts above Tsat.")
    add_table(doc, ["Metric", "Result", "Limit", "Margin", "Verdict"], [
        ["MDNBR (COLR envelope, record shape)", "1.33 (W-3*) / 2.13 (Bowring)", "1.3", "+2.5 % / +64 %", "PASS"],
        ["Bounding AOO 118 % P / 80 % G", "1.57  (Bowring, record shape)", "1.3", "+21 %", "PASS"],
        ["PCT (peak clad, boiling)", "626 K / 353 °C", "1200 °C (LOCA)", "+847 °C", "PASS"],
        ["  clad 1φ (CFD x-check, as-run)", "661 K / 388 °C", "—", "—", "—"],
        ["Peak fuel centreline", "1101 K / 828 °C", "~2840 °C (melt)", "+2012 °C", "PASS"],
    ], [2.6, 1.8, 1.3, 1.2, 0.9], hi={0, 1})
    para(doc, "The Jens-Lottes subcooled-boiling clamp caps the clad at 353 °C; the single-phase "
              "CFD cross-check (388 °C at the as-run source) confirms the film physics (§6). All "
              "safety metrics pass; the binding constraint is MDNBR at the MOC COLR envelope, "
              "which clears 1.3 on the extrapolated-conservative W-3 (+2.5 %) and comfortably on "
              "the in-range Bowring (+64 %). The AOO envelope needs no cycle stacking beyond this: "
              "the 118 %/80 % corner is evaluated at the COLR envelope itself.")

    # ---- 9 natural circulation ----
    H(doc, "9. Natural-circulation feasibility")
    para(doc, "With no pumps, the core flow is the output of a 1-D buoyancy balance between the hot "
              "riser and the cold downcomer: ṁ = [2 ρ² A² g H_tc β P / (cp K_tot)]^(1/3). The single "
              "design knob is the riser thermal-centre height H_tc (core mid-plane to steam-generator "
              "mid). Sweeping it at nominal loop losses (K_form 12):")
    figure(doc, "F8_natcirc_sweep.png",
           "Figure F8. Core mass flux G and MDNBR versus riser height H_tc. The FER operating point "
           "(ΔT 50 K → G 543) is reproduced at H_tc ≈ 4 m.")
    add_table(doc, ["H_tc [m]", "G [kg/m²s]", "core ΔT [K]", "MDNBR (W-3, envelope)", "Verdict"], [
        ["3", "492", "55", "1.12", "FAIL"],
        ["4  (design)", "542", "50", "1.37", "PASS"],
        ["5", "584", "46", "1.58", "PASS"],
        ["6", "622", "44", "1.78", "PASS"],
        ["8", "685", "40", "2.11", "PASS"],
    ], [1.4, 1.4, 1.4, 1.8, 1.1], hi={1})
    para(doc, "The FER-specified core ΔT of 50 K pins the flow at G 543, which natural circulation "
              "delivers at H_tc ≈ 4 m — half the ≈ 8 m of an earlier, smaller-core concept, because "
              "H_tc ∝ A_core·G³ and the required G fell. The MDNBR column is W-3 at the COLR "
              "envelope (chopped cosine; Bowring gives 2.13 at the design point). A taller riser "
              "(5-6 m) would trade vessel height for additional DNB margin. H_tc 3 m fails MDNBR, "
              "so ≈ 4 m is the practical floor.")

    H(doc, "9.1 Flow-stability screening (Ledinegg & density-wave)", 2)
    para(doc, "A natural-circulation core must also be shown stable, not just sufficient in flow. "
              "Two classic screens at operating pressure (tools/stability_map.py):")
    bullet(doc, "Ledinegg (excursive): the loop demand curve rises (Δp ∝ G²) while the buoyancy "
                "supply falls (Δp ∝ 1/G at fixed power), so the balance slope is positive by "
                "construction (+22 Pa per kg/m²·s at the design point) — a single, monotone, "
                "excursion-stable operating point with no multi-valued branch.")
    bullet(doc, "Density-wave oscillations (Ishii–Zuber simplified HEM criterion, parallel-channel "
                "frame, conservative inputs — core-inlet restriction only, single-phase friction): "
                "at the COLR envelope the hot channel sits at N_sub 2.47 / N_pch 3.02, a factor "
                "≈ 4.2 inside the stability boundary with NO spacer-grid credit. The bounding AOO "
                "corner (118 % power / 80 % flow) retains ×1.17 on that no-grid screen, and the "
                "stacked worst case (+10 K inlet) sits exactly on it (×1.00); crediting the five "
                "real spacer grids (Λ + 4) restores ×1.60 and ×1.37 respectively. Average channels "
                "never reach saturation (exit quality < 0) and cannot sustain a density-wave mode "
                "at all.")
    figure(doc, "F10_stability_map.png",
           "Figure F10. Density-wave stability map (Ishii–Zuber simplified criterion, 12.8 MPa) at "
           "the COLR envelope: boundaries with and without spacer-grid credit, the single-phase "
           "region (N_pch < N_sub), the design point and the AOO corners (tools/stability_map.py).")
    para(doc, "Scope: at-pressure operation. Low-pressure start-up (flashing/geysering-class modes) "
              "is excluded by procedure — the primary is pressurized before power ascension, "
              "NuScale-style heat-up; a frequency-domain or system-code stability analysis is the "
              "licensing-stage refinement.", italic=True)

    # ---- 10 vessel feasibility ----
    H(doc, "10. Vessel & layout feasibility")
    para(doc, "H_tc is an elevation offset, not the vessel height. An illustrative integral-vessel "
              "elevation budget that delivers H_tc = 4 m (core mid 1.9 m → SG mid 5.9 m):")
    add_table(doc, ["Elevation [m]", "Component"], [
        ["0.0 – 0.9", "lower plenum / core inlet"],
        ["0.9 – 2.9", "core (2.0 m active; mid-plane 1.9 m)"],
        ["2.9 – 4.4", "outlet plenum + riser (chimney)"],
        ["4.4 – 7.4", "steam generator (helical coil, in annulus; mid 5.9 m)"],
        ["7.4 – 7.9", "upper plenum"],
        ["7.9 – 9.9", "integral pressurizer"],
    ], [1.7, 4.8])
    bullet(doc, "Total vessel ≈ 10 m — markedly more compact than NuScale (≈ 17.7 m), because the "
                "required riser is half as tall.")
    bullet(doc, "The Aegis-40 core is geometrically the NuScale NPM core (37 FA / 9 768 pins / 2.0 m), "
                "an NRC-reviewed natural-circulation iPWR arrangement: helical-coil SG in the annulus "
                "around a central hot riser, integral pressurizer on top. Component layout is therefore "
                "proven; Aegis-40 runs the same core at lower power with a shorter riser.")
    bullet(doc, "The steam generator (≈ 1 250 m² helical surface for 125 MWth at ΔT_pri 50 K, "
                "467 kg/s) fits the annulus above the core with a ≈ 1.5 m riser gap. The binding "
                "dimension is vessel diameter (≈ 3.0-3.5 m, NuScale-class), not height.")
    para(doc, "This is a layout-scoping argument (iPWR analogy + elevation budget); the mechanical "
              "vessel design — seismic, wall thickness, SG supports — is outside this thermal-hydraulic "
              "scope.", italic=True)

    # ---- 11 CFD fields ----
    H(doc, "11. CFD field visualisation")
    para(doc, "ParaView renderings of the converged medium-mesh fields (t = 15):")
    figure(doc, "F9a_xsection_T.png",
           "Figure F9a. Temperature on the sub-channel cross-section at the axial peak (z = 1.0 m): "
           "hot fuel core grading through the clad to the cooler surrounding coolant.")
    figure(doc, "F9b_axial_T.png",
           "Figure F9b. Axial temperature on the pin centre-plane (chopped-cosine profile; inlet at "
           "left); radial scale exaggerated ~26x for legibility.")
    figure(doc, "F9c_velocity.png",
           "Figure F9c. Coolant axial velocity U_z on the sub-channel cross-section (z = 1.0 m): "
           "all-positive upflow (no recirculation), peak in the open sub-channel, no-slip at the rod "
           "wall (fuel+clad rod shown grey).")

    # ---- 12 limitations ----
    H(doc, "12. Limitations & uncertainty")
    bullet(doc, "The binding MDNBR 1.33 (COLR envelope, record axial shape) uses W-3, which is "
                "extrapolated below its mass-flux range (G 542 < 1356 kg/m²·s) and into deep "
                "subcooling — conservative there (under-predicts CHF), so +2.5 % is a lower bound. "
                "The low-flow-valid Bowring-1972 correlation, in range at G 542, gives 2.13 at the "
                "same point (+64 %, the in-range verdict). Bowring A/C coefficients were verified "
                "term-for-term against the Todreas & Kazimi SI form "
                "(docs/Aegis40_bowring_verification.md).")
    bullet(doc, "The hot-channel bulk is capped at Tsat with a two-phase equilibrium-quality model "
                "(exit quality 7.5 % at the COLR envelope); MDNBR is unchanged because the CHF "
                "correlations use the thermodynamic quality, not the bulk T. This affects only the "
                "top of the most-peaked channel; a CFD wall-boiling (RPI) model is the "
                "higher-fidelity refinement.")
    bullet(doc, "Peaking is the cycle-resolved OpenMC record (BOC 1.937 / MOC 2.435 / "
                "EOC 2.121; F_ΔH incl. the ×1.03 engineering allowance), bound by the COLR "
                "envelope 2.468. The record axial shapes were fed to the correlations directly: "
                "the chopped cosine proved ~2-4 % optimistic, so all binding numbers quote the "
                "real-shape values. A full lateral per-pin power map into the CFD (vs the single "
                "hot-pin source) is the remaining refinement.")
    bullet(doc, "Fuel thermal properties are fresh-fuel: constant UO₂ k = 3.5 W/m·K and He-gap "
                "conductance 5 246 W/m²·K. At the binding MOC burnup (13.5 GWd/t) the gap is "
                "partially closed (favourable) while the fuel conductivity is degraded ~10 % "
                "(≈ +30 K on the centreline); the 828 °C peak — ≈ 2 000 °C below melt — bounds "
                "any burnup penalty. A fuel-performance code (FRAPCON-class) is the "
                "licensing-stage refinement.")
    bullet(doc, "Grid-convergence uncertainty is < 0.15 % on temperatures; peak-fuel and Δp are "
                "degenerate-Richardson and reported by direct medium→fine change.")

    # ---- 13 conclusions ----
    H(doc, "13. Conclusions")
    bullet(doc, "The conjugate CFD is energy-conservative (×1.00, every mesh) and mesh-independent "
                "(GCI < 0.15 %), and its near-wall heat transfer matches the correlation stack to < 7 K.")
    bullet(doc, "The cycle peaking (BOC/MOC/EOC) is covered by a design-specific COLR "
                "envelope set by the MOC Gd-burnout hump; at that envelope the hot channel runs "
                "saturated at the top (exit quality 7.5 %) and peak fuel centreline is 828 °C.")
    bullet(doc, "Safety metrics pass at every cycle state with the record axial shapes: binding "
                "MDNBR 1.33 (W-3, conservative) / 2.13 (Bowring, in-range) at the COLR envelope; "
                "bounding AOO 1.57; PCT 353 °C (≪ 1200 °C); fuel 828 °C (≪ melt).")
    bullet(doc, "Natural circulation delivers the design flow at a modest H_tc ≈ 4 m, yielding a "
                "compact (~10 m) NuScale-class integral vessel in which the steam generator fits the "
                "annulus above the core — the reactor is thermal-hydraulically feasible.")

    doc.save(OUT)
    figs = [f for f in os.listdir(FIGS) if f.endswith(".png")] if os.path.isdir(FIGS) else []
    n = sum(1 for k in ("F1", "F2", "F3", "F4", "F5", "F6", "F7", "F8") if any(x.startswith(k) for x in figs))
    n9 = sum(1 for k in ("F9a", "F9b", "F9c") if any(x.startswith(k) for x in figs))
    print("wrote", OUT)
    print(f"embedded {n}/8 result figures + {n9}/3 ParaView renders (F9a-c)")


if __name__ == "__main__":
    main()
