#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Build the NuScale VALIDATION report (.docx) for the competition.

Verification & Validation of the OpenFOAM conjugate-pin CFD + correlation-stack
toolchain against the NuScale NPM-160 reference. Embeds F1-F7 from docs/figs/;
embeds the ParaView CFD-field figures (F8a-c, generated headless by
tools/make_paraview_figs.py). Regenerable:  python3 tools/build_report_nuscale.py
"""
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
DOCS = os.path.normpath(os.path.join(HERE, "..", "docs"))
FIGS = os.path.join(DOCS, "figs")
OUT = os.path.join(DOCS, "NuScale_validation_report.docx")

NAVY = RGBColor(0x1F, 0x4E, 0x79)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
GREY = RGBColor(0x77, 0x77, 0x77)
HDRBG, ALTBG = "1F4E79", "EAF1F7"


# ----------------------------------------------------------------- doc helpers
def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), fill)
    tcPr.append(sh)


def set_cell(cell, text, bold=False, color=None, align=None, size=9.5):
    cell.text = ""
    p = cell.paragraphs[0]
    if align:
        p.alignment = align
    r = p.add_run(text); r.bold = bold; r.font.size = Pt(size); r.font.name = "Arial"
    if color:
        r.font.color.rgb = color


def add_table(doc, headers, rows, widths, hi=None):
    hi = hi or set()
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for j, h in enumerate(headers):
        c = t.rows[0].cells[j]
        set_cell(c, h, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF)); shade(c, HDRBG)
        c.width = Inches(widths[j])
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j, val in enumerate(row):
            set_cell(cells[j], val, bold=(i in hi) or (j == 0),
                     color=GREEN if i in hi else None)
            shade(cells[j], ALTBG if (i in hi or i % 2 == 1) else "FFFFFF")
            cells[j].width = Inches(widths[j])
    for row in t.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr(); b = OxmlElement("w:tcBorders")
            for e in ("top", "left", "bottom", "right"):
                el = OxmlElement(f"w:{e}"); el.set(qn("w:val"), "single")
                el.set(qn("w:sz"), "4"); el.set(qn("w:color"), "BBBBBB"); b.append(el)
            tcPr.append(b)
    doc.add_paragraph()


def H(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        r.font.name = "Arial"; r.font.color.rgb = NAVY
    return h


def para(doc, text, italic=False, size=10.5):
    p = doc.add_paragraph(); r = p.add_run(text)
    r.font.size = Pt(size); r.font.name = "Arial"; r.italic = italic
    return p


def bullet(doc, text):
    doc.add_paragraph(text, style="List Bullet")


def caption(doc, text):
    p = doc.add_paragraph(); p.alignment = AL.CENTER
    r = p.add_run(text); r.italic = True; r.font.size = Pt(9); r.font.name = "Arial"
    r.font.color.rgb = GREY


def figure(doc, fname, cap, width=6.2):
    path = os.path.join(FIGS, fname)
    if os.path.exists(path):
        doc.add_picture(path, width=Inches(width))
        doc.paragraphs[-1].alignment = AL.CENTER
        caption(doc, cap)
    else:                                   # labelled placeholder
        t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = t.rows[0].cells[0]; c.width = Inches(width)
        set_cell(c, f"[ {fname} — ParaView screenshot to be inserted ]",
                 align=AL.CENTER, color=GREY)
        shade(c, "F2F2F2")
        tcPr = c._tc.get_or_add_tcPr(); b = OxmlElement("w:tcBorders")
        for e in ("top", "left", "bottom", "right"):
            el = OxmlElement(f"w:{e}"); el.set(qn("w:val"), "dashed")
            el.set(qn("w:sz"), "6"); el.set(qn("w:color"), "999999"); b.append(el)
        tcPr.append(b)
        for _ in range(4):
            c.add_paragraph()
        caption(doc, cap)


# ----------------------------------------------------------------- the report
def main():
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10.5)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for m in ("top", "bottom", "left", "right"):
        setattr(sec, f"{m}_margin", Inches(0.9))

    # ---- title ----
    p = doc.add_paragraph(); p.alignment = AL.CENTER
    r = p.add_run("Conjugate Fuel-Pin CFD — Verification & Validation")
    r.bold = True; r.font.size = Pt(20); r.font.name = "Arial"; r.font.color.rgb = NAVY
    p = doc.add_paragraph(); p.alignment = AL.CENTER
    r = p.add_run("OpenFOAM chtMultiRegionFoam toolchain validated against the "
                  "NuScale NPM-160 reference")
    r.font.size = Pt(12); r.font.name = "Arial"; r.font.color.rgb = RGBColor(0x55,0x55,0x55)
    p = doc.add_paragraph(); p.alignment = AL.CENTER
    r = p.add_run("Thermal-hydraulics validation case  |  June 2026")
    r.italic = True; r.font.size = Pt(10); r.font.color.rgb = GREY
    doc.add_paragraph()

    # ---- 1 objective ----
    H(doc, "1. Objective & scope")
    para(doc, "This report validates the OpenFOAM conjugate heat-transfer toolchain — a "
              "3-region fuel/clad/coolant CFD model (chtMultiRegionFoam) coupled to a "
              "correlation-based safety post-processor — against the NuScale NPM-160 "
              "integral PWR. NuScale is a published, well-characterised reference of the "
              "same class as the competition design (≤40 MWe natural-circulation iPWR), so "
              "reproducing its operating point demonstrates the toolchain is correct before "
              "it is applied to the final design geometry.")
    para(doc, "The verification & validation follows a conservation-first protocol: energy "
              "balance is the first acceptance test on every mesh, before any temperature is "
              "trusted.")

    # ---- 2 reference point ----
    H(doc, "2. Reference operating point (NuScale NPM-160)")
    add_table(doc, ["Quantity", "Value", "Source"], [
        ["Core thermal power", "160 MWt", "ASTEC model, Front. Energy Res. 10:1036142"],
        ["System pressure", "12.8 MPa", "—"],
        ["Core inlet / hot-leg T", "531 K / 583 K (258 / 310 °C)", "—"],
        ["Primary mass flow", "587 kg/s", "—"],
        ["Fuel pins", "9 768  (37 FA × 264, 17×17)", "—"],
        ["Active length / pitch", "2.0 m / 12.6 mm", "—"],
        ["Pellet / clad OD", "8.19 mm / 9.5 mm  (UO₂ / Zircaloy-4)", "—"],
        ["Total peaking F_q", "1.923  (BOC PPF)", "NuScale neutronic study, S0306454919304979"],
        ["Hot-pin peak linear power", "15.75 kW/m  (q‴ 2.99e8 W/m³)", "= q′_avg 8.19 × F_q"],
    ], [2.1, 2.5, 2.4])

    # ---- 3 methodology ----
    H(doc, "3. Methodology")
    H(doc, "3.1 Conjugate CFD", 2)
    bullet(doc, "Solver chtMultiRegionFoam (OpenFOAM v2412); 3 regions fuel/clad/coolant on "
                "a body-fitted butterfly O-grid; turbulence k-ω SST.")
    bullet(doc, "Axial power profile: chopped cosine q‴(z)=q_peak·cos(π(z−L/2)/Lₑ), Lₑ=1.2 L, "
                "imposed as a runtime scalarCodedSource in the fuel region.")
    bullet(doc, "Fuel-clad gap modelled as a contact resistance (thicknessLayers / "
                "kappaLayers, applied symmetrically on both coupling patches).")
    H(doc, "3.2 Safety post-processor (correlation stack)", 2)
    para(doc, "MDNBR and PCT are obtained from a validated correlation stack — the "
              "COBRA/VIPRE/CTF subchannel-code approach — fed the CFD-trusted axial power "
              "shape, coolant bulk T(z) and flow:")
    bullet(doc, "MDNBR: W-3 critical-heat-flux correlation with the Tong non-uniform-flux "
                "F-factor.")
    bullet(doc, "PCT / fuel centreline: Dittus-Boelter film coefficient + 1-D radial "
                "conduction stack (film → clad → He-gap → fuel), with a Jens-Lottes "
                "subcooled-boiling wall clamp.")
    para(doc, "The correlation stack is used because it is the licensed-code industry "
              "standard and because it incorporates the CHF correlation and subcooled-boiling "
              "physics that a single-phase CFD cannot — NOT because the CFD is inaccurate. As "
              "Section 5 shows, the CFD is energy-conservative and its near-wall heat transfer "
              "agrees with Dittus-Boelter to ~1 %.", italic=True)

    # ---- 4 mesh & numerics ----
    H(doc, "4. Mesh & numerics")
    add_table(doc, ["Mesh", "Cells (fuel/clad/coolant)", "Coolant AR", "Refinement r"], [
        ["coarse", "35 700  (16.1 / 8.4 / 11.2 k)", "≈ 38", "—"],
        ["medium (production)", "103 600  (47.6 / 22.4 / 33.6 k)", "≈ 57", "r₃₂ = 1.426"],
        ["fine", "279 888  (126.2 / 65.9 / 87.8 k)", "≈ 80", "r₂₁ = 1.393"],
    ], [1.7, 2.7, 1.1, 1.5], hi={1})
    bullet(doc, "Schemes: div(phi,h) Gauss upwind (conservation gate) → limitedLinear "
                "(accuracy); nNonOrthogonalCorrectors 2; nOuterCorrectors 2.")
    bullet(doc, "Linear solvers GAMG (p_rgh) / smoothSolver; relTol 0.01, no maxIter caps; "
                "real cp = 5250 J/kg·K.")

    # ---- 5 verification ----
    H(doc, "5. Verification")
    H(doc, "5.1 Grid convergence (ASME V&V-20 GCI)", 2)
    figure(doc, "F1_grid_convergence.png",
           "Figure F1. Grid convergence of outlet T, peak clad T and peak fuel T over the "
           "three meshes, with Richardson extrapolation (h→0) and the fine-grid GCI.")
    add_table(doc, ["Metric", "coarse", "medium", "fine", "GCI_fine"], [
        ["Outlet T  [K]", "603.60", "604.16", "604.43", "0.06 %"],
        ["Peak clad T  [K]", "660.03", "660.43", "660.63", "0.05 %"],
        ["Peak fuel T  [K]", "1108.4", "1110.6", "1112.7", "0.19 % (direct)"],
    ], [2.0, 1.3, 1.3, 1.3, 1.5], hi={0, 1})
    para(doc, "The medium↔fine change is below 0.2 % on every metric and the fine-grid GCI is "
              "sub-0.1 % on the convergent metrics, so the solution is mesh-independent. The "
              "medium (104 k) mesh is used for production.")

    H(doc, "5.2 Energy conservation (GATE-1)", 2)
    figure(doc, "F2_energy_conservation.png",
           "Figure F2. Advected power ṁ·cp·ΔT versus the design source on each mesh; "
           "conservation closes to ×1.00 (≤0.8 %) on every grid.")
    para(doc, "The advected power equals the imposed source (and the coolant_to_clad wall-flux "
              "integral) to within 0.8 % on every mesh — the model is energy-conservative, the "
              "prerequisite for trusting any temperature.")

    H(doc, "5.3 Power-linearity check", 2)
    para(doc, "Because the single-phase conjugate problem with constant properties is linear in "
              "the heat source, the medium mesh was additionally re-run natively at the "
              "NuScale-calibrated power (F_q 1.923). The native fields reproduce the exact "
              "power-rescale of the F_q 2.17 run to ±0.01 K / ±0 Pa, confirming the linearity "
              "used to present the mesh study at the design point.")

    # ---- 6 validation ----
    H(doc, "6. Validation against NuScale")
    figure(doc, "F3_vs_nuscale.png",
           "Figure F3. CFD + correlation-stack results versus published NuScale NPM-160 data.")
    add_table(doc, ["Quantity", "NuScale (publ.)", "This work", "Δ"], [
        ["Core inlet T", "258 °C", "258 °C", "exact (BC)"],
        ["Core-avg outlet (hot leg)", "310 °C", "310 °C", "match"],
        ["Hot-channel outlet (≈ Tsat)", "331 °C", "331 °C", "match"],
        ["Peak clad T (max)", "~360 °C", "352 °C", "−8 °C"],
    ], [2.4, 1.6, 1.4, 1.3], hi={3})
    para(doc, "The core-average-equivalent outlet matches the NuScale hot leg, the hot channel "
              "sits right at the saturation boundary (subcooled boiling onset, as in the NuScale "
              "design), and the peak clad temperature agrees with the published value within "
              "~8 °C — all with no tuning.")

    # ---- 7 hot-channel ----
    H(doc, "7. Hot-channel thermal-hydraulics")
    figure(doc, "F4_axial_profiles.png",
           "Figure F4. Hot-channel axial profiles: wall heat flux q″(z); coolant bulk, clad "
           "outer (with subcooled-boiling clamp) and fuel centreline temperatures.")
    figure(doc, "F5_radial_stack.png",
           "Figure F5. Radial temperature stack at the hot spot — the conduction drops across "
           "fuel, the He-gap contact resistance, clad and the convective film.")

    # ---- 8 safety ----
    H(doc, "8. Safety margins")
    H(doc, "8.1 MDNBR (departure from nucleate boiling)", 2)
    figure(doc, "F6_dnbr_profile.png",
           "Figure F6. Hot-channel DNBR axial profile (W-3 + Tong F-factor); the minimum is the "
           "MDNBR.")
    H(doc, "8.2 PCT & subcooled boiling", 2)
    figure(doc, "F7_subcooled_boiling.png",
           "Figure F7. Clad outer temperature: single-phase film versus the subcooled-boiling "
           "(Jens-Lottes) clamp — single-phase over-predicts above Tsat.")
    add_table(doc, ["Metric", "Result", "Limit", "Margin", "Verdict"], [
        ["MDNBR", "1.977  @ z = 0.57 m", "1.3", "+52 %", "PASS"],
        ["PCT (peak clad, boiling)", "625.6 K / 352 °C", "1200 °C (LOCA)", "+848 °C", "PASS"],
        ["  clad, single-phase (CFD x-check)", "660 K / 387 °C", "—", "—", "—"],
        ["Peak fuel centreline", "1092 K / 819 °C", "~2840 °C (melt)", "+2021 °C", "PASS"],
    ], [2.3, 1.6, 1.4, 1.0, 0.9], hi={0, 1})
    para(doc, "Cross-check: the correlation stack's single-phase clad (380 °C) and fuel centreline "
              "(1112 K) agree with the CFD (387 °C, 1111 K) to within 8 K and 1 K respectively, "
              "confirming the CFD near-wall heat transfer (Dittus-Boelter ≈ CFD).")

    # ---- 9 CFD fields (placeholders) ----
    H(doc, "9. CFD field visualisation")
    para(doc, "The following ParaView renderings of the converged medium-mesh fields (t = 15) "
              "are generated headless by tools/make_paraview_figs.py (pvbatch).")
    figure(doc, "F8a_xsection_T.png",
           "Figure F8a. Temperature on a cross-section at the axial peak (z = 1.0 m): hot fuel "
           "core, clad ring, cooler coolant annulus.")
    figure(doc, "F8b_axial_T.png",
           "Figure F8b. Axial temperature slice along the pin (chopped-cosine source; radial "
           "scale exaggerated \u00d726 for legibility).")
    figure(doc, "F8c_velocity.png",
           "Figure F8c. Coolant axial velocity field — developed, all-positive upflow (no "
           "recirculation → stable throughflow).")

    # ---- 10 limitations ----
    H(doc, "10. Limitations & uncertainty")
    bullet(doc, "MDNBR uses the W-3 correlation, which is extrapolated below its mass-flux range "
                "(G 684 < 1356 kg/m²·s) and into deep subcooling (x < −0.15); in this regime W-3 "
                "is conservative (under-predicts CHF), so the +52 % margin is a lower bound. A "
                "low-flow CHF correlation is recommended for a final safety case.")
    bullet(doc, "The 1-D single-phase energy balance lets the hot-channel bulk slightly exceed "
                "Tsat (~1.8 K) at the outlet; a two-phase enthalpy model would cap it at Tsat. "
                "This affects only the last few cm of the most-peaked channel.")
    bullet(doc, "Grid-convergence uncertainty (GCI) is < 0.1 % on outlet/clad temperatures; the "
                "peak-fuel and Δp metrics are degenerate-Richardson and are reported by their "
                "direct medium→fine change (< 0.2 %, < 1.6 %).")

    # ---- 11 conclusions ----
    H(doc, "11. Conclusions")
    bullet(doc, "The conjugate CFD is energy-conservative (×1.00, every mesh) and "
                "mesh-independent (GCI < 0.1 %).")
    bullet(doc, "It reproduces the NuScale operating point — hot-leg temperature, the "
                "saturation-boundary hot channel and the peak clad temperature — with no tuning.")
    bullet(doc, "Safety metrics pass with margin: MDNBR 1.977 (> 1.3) and PCT 352 °C (≪ 1200 °C).")
    bullet(doc, "The toolchain is validated and ready to be applied to the competition design "
                "geometry (edit parameters, re-mesh, re-run the same scripts).")

    doc.save(OUT)
    print("wrote", OUT)
    figs = [f for f in os.listdir(FIGS) if f.endswith(".png")] if os.path.isdir(FIGS) else []
    print(f"embedded {sum(1 for n in ('F1','F2','F3','F4','F5','F6','F7') if any(x.startswith(n) for x in figs))}/7 result figures + "
          f"{sum(1 for n in ('F8a','F8b','F8c') if any(x.startswith(n) for x in figs))}/3 ParaView renders (F8a-c)")


if __name__ == "__main__":
    main()
