import sys
from docx import Document
p = r"D:\projects\teknofest-2026-aegis-40-ipwr\docs\competition\fer\Aegis40_TH_report (2).docx"
d = Document(p)
print(f"paragraphs={len(d.paragraphs)} tables={len(d.tables)}")
print("=" * 60, "HEADINGS + TEXT", "=" * 60)
for para in d.paragraphs:
    t = para.text.strip()
    if not t:
        continue
    sty = para.style.name if para.style else ""
    tag = f"[{sty}] " if sty.lower().startswith(("heading", "title")) else ""
    print(tag + t)
for ti, tab in enumerate(d.tables):
    print(f"\n----- TABLE {ti} -----")
    for row in tab.rows:
        print(" | ".join(c.text.strip() for c in row.cells))
